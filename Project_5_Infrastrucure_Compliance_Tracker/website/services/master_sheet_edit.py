"""Safe, audited edits for master-workbook rows without an imported snapshot row."""
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from website.models import AuditLog, MasterWorkbook
from website.services.errors import ImportValidationError
from website.services.excel_row_edit import _atomic_replace, _cell_value_from_text
from website.services.master_workbook import (
    _headers,
    import_master_workbook,
    parameter_sheet_name,
    preview_master_workbook,
)


MANAGED_HEADINGS = {
    "lastupdated", "lastupdatedon", "lastmodified", "checkdate",
    "changesmadeby", "changedby", "updatedby",
}


def _key(heading):
    return "".join(character for character in heading.casefold() if character.isalnum())


def edit_master_sheet_row(parameter_name, source_row, updates, actor):
    """Write an empty-tab row to the owner-selected workbook and refresh its snapshots."""
    with transaction.atomic():
        config = MasterWorkbook.objects.select_for_update().first()
        if not config or not config.path:
            raise ImportValidationError("The owner has not configured the master workbook write-back path.")
        target = Path(config.path)
        try:
            original = target.read_bytes()
        except OSError as exc:
            raise ImportValidationError("The configured master workbook is not available on this computer.") from exc
        if sha256(original).hexdigest() != config.content_sha256:
            raise ImportValidationError("The master workbook changed outside the web app. Re-import it before editing.")
        try:
            workbook = load_workbook(BytesIO(original), data_only=False)
        except Exception as exc:
            raise ImportValidationError("The configured master workbook could not be opened for editing.") from exc
        try:
            sheet_name = parameter_sheet_name(parameter_name)
            if sheet_name not in workbook.sheetnames or "Main Data" not in workbook.sheetnames:
                raise ImportValidationError(f"The master workbook has no editable worksheet for {parameter_name}.")
            sheet = workbook[sheet_name]
            inventory = workbook["Main Data"]
            if source_row < 2 or source_row > max(sheet.max_row, inventory.max_row):
                raise ImportValidationError("That worksheet row is no longer available. Reload the spreadsheet.")
            headings = _headers(sheet)
            inventory_headings = _headers(inventory)
            sheet_indexes = {heading: index + 1 for index, heading in enumerate(headings)}
            inventory_indexes = {heading: index + 1 for index, heading in enumerate(inventory_headings)}
            allowed = set(headings) | {heading for heading in ("Hostname", "IP Address") if heading in inventory_indexes}
            unexpected = set(updates) - allowed
            if unexpected:
                raise ImportValidationError("One or more submitted columns are no longer present in the worksheet.")

            now = timezone.localtime().replace(tzinfo=None, second=0, microsecond=0)
            actual_updates = dict(updates)
            for heading in headings:
                if _key(heading) in {"lastupdated", "lastupdatedon", "lastmodified", "checkdate"}:
                    actual_updates[heading] = now
                elif _key(heading) in {"changesmadeby", "changedby", "updatedby"}:
                    actual_updates[heading] = actor
            written = {}
            for heading, text in actual_updates.items():
                if heading in {"Hostname", "IP Address"}:
                    target_sheet = inventory
                    column = inventory_indexes[heading]
                else:
                    target_sheet = sheet
                    column = sheet_indexes[heading]
                cell = target_sheet.cell(row=source_row, column=column)
                if cell.data_type == "f":
                    raise ImportValidationError(f"{heading} is formula-driven and cannot be changed here.")
                value = text if isinstance(text, datetime) else _cell_value_from_text(text, cell.value, heading)
                if isinstance(value, str) and value.lstrip().startswith("="):
                    raise ImportValidationError(f"Formula-like values are not allowed in {heading}. Enter plain text only.")
                if heading == "Hostname" and not str(value or "").strip():
                    raise ImportValidationError("Hostname cannot be blank because it identifies the server.")
                cell.value = value
                if isinstance(value, datetime) and cell.number_format == "General":
                    cell.number_format = "yyyy-mm-dd hh:mm"
                written[heading] = value
            if getattr(workbook, "calculation", None):
                workbook.calculation.calcMode = "auto"
                workbook.calculation.fullCalcOnLoad = True
                workbook.calculation.forceFullCalc = True
            output = BytesIO()
            workbook.save(output)
            updated = output.getvalue()
        except ImportValidationError:
            raise
        except Exception as exc:
            raise ImportValidationError("The worksheet row could not be updated safely.") from exc
        finally:
            workbook.close()

        # Validate the changed package before replacing the user's workbook.
        preview_master_workbook(updated)
        if sha256(target.read_bytes()).hexdigest() != config.content_sha256:
            raise ImportValidationError("The master workbook changed while this row was being edited. Reload and try again.")
        backup = target.with_name(f"{target.stem}.before_web_edits{target.suffix}")
        if not backup.exists():
            with backup.open("xb") as handle:
                handle.write(original)
                handle.flush()
        _atomic_replace(target, updated)
        try:
            import_master_workbook(
                updated,
                original_filename=config.original_filename,
                target_path=str(target),
            )
        except Exception:
            _atomic_replace(target, original)
            raise
        detail = ", ".join(sorted(written))
        AuditLog.objects.create(
            action="master_sheet_row_edited",
            detail=f"{parameter_name}, Excel row {source_row}, edited by {actor}. Updated: {detail}.",
        )
    return written
