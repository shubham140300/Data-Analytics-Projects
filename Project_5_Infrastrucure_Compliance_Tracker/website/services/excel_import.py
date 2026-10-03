"""Read and validate Excel files into normalized source-row dictionaries."""
from collections import defaultdict
from calendar import month_abbr, month_name
from datetime import date, datetime
from io import BytesIO
import json
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.xml import DEFUSEDXML

from website.models import ComplianceStatus
from website.services.errors import ImportValidationError


REQUIRED_COLUMNS = ("Host Name", "IP Address", "OPERATING SYSTEM", "Environment")
STANDARD_STATUS_MAP = {
    "Compliant": ComplianceStatus.COMPLIANT,
    "Non-Compliant": ComplianceStatus.NON_COMPLIANT,
    "Not Applicable": ComplianceStatus.NOT_APPLICABLE,
}
MAX_PACKAGE_BYTES = 250 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2000
MAX_WORKSHEET_ROWS = 100_000
MAX_WORKSHEET_COLUMNS = 200
MAX_WORKSHEET_CELLS = 2_000_000
ACTIVE_CONTENT_PARTS = (
    "xl/vbaproject.bin",
    "xl/externallinks/",
    "xl/embeddings/",
    "xl/activex/",
    "xl/connections.xml",
)


def _cell_for_json(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _text(value):
    return "" if value is None else str(value).strip()


def _date_from_heading(heading):
    match = re.search(r"\b(\d{1,2})[\s,_-]+([A-Za-z]+)[\s,_-]+(20\d{2})\b", heading)
    if not match:
        return None
    month_lookup = {}
    for number in range(1, 13):
        month_lookup[month_name[number].casefold()] = number
        month_lookup[month_abbr[number].casefold()] = number
    month = month_lookup.get(match.group(2).casefold())
    if month is None:
        return None
    try:
        return date(int(match.group(3)), month, int(match.group(1)))
    except ValueError:
        return None


def _choose_status_column(headers, data_rows, requested_column=""):
    if requested_column:
        if requested_column not in headers:
            raise ImportValidationError(
                f"The status heading '{requested_column}' was not found. Choose one of: "
                + ", ".join(headers)
            )
        if requested_column in REQUIRED_COLUMNS:
            raise ImportValidationError("The selected status column is also a required server detail column.")
        return requested_column

    candidates = []
    for index, heading in enumerate(headers):
        if heading in REQUIRED_COLUMNS:
            continue
        sample_values = {_text(row[index]) for row in data_rows[:200] if index < len(row)}
        if sample_values.intersection(STANDARD_STATUS_MAP):
            candidates.append(heading)
    if len(candidates) == 1:
        return candidates[0]
    remaining = [heading for heading in headers if heading not in REQUIRED_COLUMNS]
    if not candidates and len(remaining) == 1:
        return remaining[0]
    if not candidates:
        raise ImportValidationError(
            "I could not identify one status column. Enter the exact status column heading from the workbook."
        )
    raise ImportValidationError(
        "Several columns look like status columns. Enter the exact heading for the one to import: "
        + ", ".join(candidates)
    )


def preview_workbook(content, requested_status_column=""):
    """Read workbook structure and rows without changing the source workbook."""
    if not zipfile.is_zipfile(BytesIO(content)):
        raise ImportValidationError("This file is not a readable .xlsx workbook.")
    if not DEFUSEDXML:
        raise ImportValidationError("The secure Excel XML parser is missing. Restart the app after installing its requirements.")
    try:
        with zipfile.ZipFile(BytesIO(content), "r") as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ImportValidationError("The workbook contains too many internal files to process safely.")
            if sum(item.file_size for item in members) > MAX_PACKAGE_BYTES:
                raise ImportValidationError("The expanded workbook is too large to process safely.")
            if any(item.flag_bits & 0x1 for item in members):
                raise ImportValidationError("Encrypted workbooks are not supported.")
            names = [item.filename.casefold() for item in members]
            if "xl/workbook.xml" not in names:
                raise ImportValidationError("This ZIP file does not contain a valid Excel workbook.")
            if any(name.endswith(part) or name.startswith(part) for name in names for part in ACTIVE_CONTENT_PARTS):
                raise ImportValidationError("Workbooks with macros, external links, embedded files, or active controls are not supported.")
    except ImportValidationError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise ImportValidationError("The workbook package could not be inspected safely.") from exc
    try:
        formula_check = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            total_cells = 0
            for worksheet in formula_check.worksheets:
                rows = worksheet.max_row or 0
                columns = worksheet.max_column or 0
                total_cells += rows * columns
                if rows > MAX_WORKSHEET_ROWS or columns > MAX_WORKSHEET_COLUMNS or total_cells > MAX_WORKSHEET_CELLS:
                    raise ImportValidationError(
                        f"The workbook is larger than the safe limit of {MAX_WORKSHEET_ROWS:,} rows, "
                        f"{MAX_WORKSHEET_COLUMNS} columns, and {MAX_WORKSHEET_CELLS:,} total cells."
                    )
                if any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row):
                    raise ImportValidationError("Formula cells are not accepted. Upload a values-only .xlsx workbook.")
        finally:
            formula_check.close()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        if isinstance(exc, ImportValidationError):
            raise
        raise ImportValidationError("The workbook could not be opened. Save it as an .xlsx file and try again.") from exc

    try:
        sheet = next((item for item in workbook.worksheets if item.max_row), None)
        if sheet is None:
            raise ImportValidationError("The workbook does not contain a worksheet with data.")
        if sheet.max_row > MAX_WORKSHEET_ROWS or sheet.max_column > MAX_WORKSHEET_COLUMNS or sheet.max_row * sheet.max_column > MAX_WORKSHEET_CELLS:
            raise ImportValidationError(
                f"The first worksheet is larger than the safe limit of {MAX_WORKSHEET_ROWS:,} rows and {MAX_WORKSHEET_COLUMNS} columns."
            )
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if raw_headers is None:
            raise ImportValidationError("The first worksheet is empty.")
        headers = ["" if item is None else str(item) for item in raw_headers]
        while headers and not headers[-1].strip():
            headers.pop()
        if not headers or any(not item.strip() for item in headers):
            raise ImportValidationError("Every column in the header row needs a heading.")
        if len(set(headers)) != len(headers):
            raise ImportValidationError("The header row contains repeated column names. Rename the duplicate headings and try again.")

        missing = [name for name in REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ImportValidationError("Required columns were not found: " + ", ".join(missing) + ".")

        raw_data = []
        for source_row, raw_values in enumerate(iterator, start=2):
            values = list(raw_values[:len(headers)])
            values.extend([None] * (len(headers) - len(values)))
            if not any(value is not None and str(value).strip() for value in values):
                continue
            raw_data.append((source_row, values))
        if not raw_data:
            raise ImportValidationError("The first worksheet has headings but no data rows.")

        status_column = _choose_status_column(headers, [row for _, row in raw_data], requested_status_column)
        column_indexes = {name: headers.index(name) for name in headers}
        rows = []
        seen_rows = set()
        status_counts = defaultdict(int)
        groups = defaultdict(list)
        missing_hostname_rows = 0
        missing_ip_rows = 0
        duplicate_rows = 0
        for source_row, values in raw_data:
            source_data = {
                heading: _cell_for_json(values[index])
                for index, heading in enumerate(headers)
            }
            signature = json.dumps(source_data, ensure_ascii=False, sort_keys=True, default=str)
            is_duplicate = signature in seen_rows
            if is_duplicate:
                duplicate_rows += 1
            seen_rows.add(signature)

            hostname = _text(values[column_indexes["Host Name"]])
            ip_address = _text(values[column_indexes["IP Address"]])
            operating_system = _text(values[column_indexes["OPERATING SYSTEM"]])
            environment = _text(values[column_indexes["Environment"]])
            status_cell = values[column_indexes[status_column]]
            raw_status = "" if status_cell is None else str(status_cell)
            status_key = raw_status.strip()
            mapped_status = STANDARD_STATUS_MAP.get(status_key, ComplianceStatus.UNMAPPED)
            status_counts[status_key or "(blank)"] += 1
            normalized_hostname = hostname.casefold() if hostname else ""
            issue_notes = []
            if mapped_status == ComplianceStatus.UNMAPPED:
                issue_notes.append("Unmapped source status label; current state remains Unmapped.")
            if not hostname:
                missing_hostname_rows += 1
                issue_notes.append("Missing Host Name; this row is retained but cannot be linked to a server.")
            if not ip_address:
                missing_ip_rows += 1
                issue_notes.append("Missing IP Address; the source row is retained.")
            issue = " ".join(issue_notes)
            if normalized_hostname:
                groups[normalized_hostname].append({
                    "hostname": hostname,
                    "ip_address": ip_address,
                    "operating_system": operating_system,
                    "environment": environment,
                    "raw_status": raw_status,
                    "mapped_status": mapped_status,
                    "source_row": source_row,
                    "issue": issue,
                })
            rows.append({
                "source_row": source_row,
                "hostname": hostname,
                "normalized_hostname": normalized_hostname,
                "ip_address": ip_address,
                "operating_system": operating_system,
                "environment": environment,
                "raw_status": raw_status,
                "mapped_status": mapped_status,
                "source_data": source_data,
                "duplicate_row": is_duplicate,
                "issue": issue,
            })

        conflicting_hosts = {
            key for key, group in groups.items()
            if len({(record["mapped_status"], record["raw_status"]) for record in group}) > 1
        }
        for row in rows:
            if row["normalized_hostname"] in conflicting_hosts:
                conflict_note = "Conflicting statuses for this Host Name in one snapshot; current state was not changed."
                row["issue"] = (row["issue"] + " " if row["issue"] else "") + conflict_note

        metadata_groups = defaultdict(list)
        for row in rows:
            if row["normalized_hostname"]:
                metadata_groups[row["normalized_hostname"]].append(row)
        multiple_ip_groups = sum(
            len({row["ip_address"] for row in group if row["ip_address"]}) > 1
            for group in metadata_groups.values()
        )

        suggested_date = _date_from_heading(status_column)
        preview = {
            "source_sheet": sheet.title,
            "headers": headers,
            "status_column": status_column,
            "suggested_date": suggested_date,
            "rows": rows,
            "status_counts": dict(status_counts),
            "total_rows": len(rows),
            "unique_servers": len(groups),
            "missing_hostname_rows": missing_hostname_rows,
            "missing_ip_rows": missing_ip_rows,
            "duplicate_rows": duplicate_rows,
            "conflicting_server_groups": len(conflicting_hosts),
            "conflicting_hostnames": conflicting_hosts,
            "multiple_ip_groups": multiple_ip_groups,
            "unmapped_status_rows": sum(
                1 for row in rows if row["mapped_status"] == ComplianceStatus.UNMAPPED
            ),
        }
        return preview
    finally:
        workbook.close()
