"""Apply an audited web edit to an imported Excel row and its workbook copy."""
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import tempfile
import uuid

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from website.models import (
    AuditLog,
    ComplianceState,
    ComplianceStatus,
    ImportBatch,
    ImportRecord,
    MasterWorkbook,
    Server,
    StatusEvent,
)
from website.services.errors import ImportValidationError
from website.services.excel_import import STANDARD_STATUS_MAP


def validate_sync_target_path(path, *, expected_hash, source_sheet, source_columns):
    """Verify an optional local workbook target matches the imported snapshot."""
    path = (path or "").strip()
    if not path:
        return ""
    target = Path(path).expanduser()
    if not target.is_absolute():
        raise ImportValidationError("Enter a full path to the original .xlsx workbook.")
    if target.suffix.casefold() != ".xlsx" or not target.is_file():
        raise ImportValidationError("The workbook path must point to an existing .xlsx file on this computer.")
    media_root = Path(settings.MEDIA_ROOT).resolve()
    target_resolved = target.resolve()
    if target_resolved == media_root or media_root in target_resolved.parents:
        raise ImportValidationError("The original workbook path must be outside the app's managed uploads folder.")
    try:
        content = target.read_bytes()
        if sha256(content).hexdigest() != expected_hash:
            raise ImportValidationError(
                "The selected workbook does not match this import. Choose the exact file that was uploaded."
            )
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The selected workbook could not be opened.") from exc
    try:
        if source_sheet not in workbook.sheetnames:
            raise ImportValidationError("The selected workbook does not contain the imported worksheet.")
        sheet = workbook[source_sheet]
        headers = ["" if value is None else str(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        while headers and not headers[-1].strip():
            headers.pop()
        if headers != source_columns:
            raise ImportValidationError("The selected workbook headings have changed. Import the latest workbook first.")
    finally:
        workbook.close()
    return str(target.resolve())


def _read_uploaded_copy(batch):
    workbook_file = batch.working_file or batch.source_file
    expected_hash = batch.working_sha256 or batch.file_sha256
    if not workbook_file:
        raise ImportValidationError("This import has no saved workbook copy to edit.")
    with workbook_file.open("rb") as source:
        content = source.read()
    if sha256(content).hexdigest() != expected_hash:
        raise ImportValidationError(
            "The saved workbook copy changed outside the web app. Re-import the latest workbook before editing."
        )
    return content


def _cell_value_from_text(text, original, heading):
    if text == "":
        return None
    if isinstance(original, datetime):
        try:
            return datetime.fromisoformat(text)
        except ValueError as exc:
            raise ImportValidationError(f"Enter {heading} as a date and time in YYYY-MM-DD HH:MM format.") from exc
    if isinstance(original, date):
        try:
            return date.fromisoformat(text)
        except ValueError as exc:
            raise ImportValidationError(f"Enter {heading} as a date in YYYY-MM-DD format.") from exc
    if isinstance(original, bool):
        if text.casefold() not in {"true", "false"}:
            raise ImportValidationError(f"Enter {heading} as True or False.")
        return text.casefold() == "true"
    if isinstance(original, int) and not isinstance(original, bool):
        try:
            return int(text)
        except ValueError as exc:
            raise ImportValidationError(f"Enter a whole number for {heading}.") from exc
    if isinstance(original, float):
        try:
            return float(text)
        except ValueError as exc:
            raise ImportValidationError(f"Enter a number for {heading}.") from exc
    return text


def _rewrite_workbook(content, record, updates):
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise ImportValidationError("The saved .xlsx workbook could not be opened for editing.") from exc
    try:
        if record.batch.source_sheet not in workbook.sheetnames:
            raise ImportValidationError("The worksheet for this imported row is missing from the saved workbook.")
        sheet = workbook[record.batch.source_sheet]
        headers = ["" if cell.value is None else str(cell.value) for cell in sheet[1]]
        while headers and not headers[-1].strip():
            headers.pop()
        if headers != record.batch.source_columns:
            raise ImportValidationError("The saved workbook headings no longer match this import.")
        if record.source_row < 2 or record.source_row > sheet.max_row:
            raise ImportValidationError("The Excel row number for this record is no longer valid.")
        header_indexes = {header: index + 1 for index, header in enumerate(headers)}
        updated_values = {}
        for heading, text in updates.items():
            column = header_indexes.get(heading)
            if column is None:
                raise ImportValidationError(f"The heading '{heading}' is not present in this workbook.")
            cell = sheet.cell(row=record.source_row, column=column)
            if cell.data_type == "f":
                raise ImportValidationError(f"{heading} is a formula cell and cannot be changed from this page.")
            value = _cell_value_from_text(text, cell.value, heading) if isinstance(text, str) else text
            if isinstance(value, str) and value.lstrip().startswith("="):
                raise ImportValidationError(f"Formula-like values are not allowed in {heading}. Enter plain text only.")
            cell.value = value
            if isinstance(value, datetime) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd"
            updated_values[heading] = value
        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), updated_values
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The workbook row could not be updated safely.") from exc
    finally:
        workbook.close()


def _rewrite_master_workbook(content, record, updates):
    """Patch a parameter tab while routing formula-linked identity fields to Main Data."""
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise ImportValidationError("The current master .xlsx workbook could not be opened for editing.") from exc
    try:
        if record.batch.source_sheet not in workbook.sheetnames or "Main Data" not in workbook.sheetnames:
            raise ImportValidationError("The master workbook is missing a required worksheet.")
        sheet = workbook[record.batch.source_sheet]
        headers = ["" if cell.value is None else str(cell.value).strip() for cell in sheet[1]]
        while headers and not headers[-1]:
            headers.pop()
        if headers != record.batch.source_columns:
            raise ImportValidationError("The master workbook headings changed. Import the latest workbook first.")
        if record.source_row < 2 or record.source_row > sheet.max_row:
            raise ImportValidationError("The Excel row number for this record is no longer valid.")

        inventory = workbook["Main Data"]
        inventory_headers = ["" if cell.value is None else str(cell.value).strip() for cell in inventory[1]]
        if "Hostname" not in inventory_headers or "IP Address" not in inventory_headers:
            raise ImportValidationError("The master inventory needs Hostname and IP Address headings.")
        row = record.source_row
        host_cell = inventory.cell(row=row, column=inventory_headers.index("Hostname") + 1)
        current_host = "" if host_cell.value is None else str(host_cell.value).strip()
        if current_host.casefold() != record.hostname.strip().casefold():
            raise ImportValidationError("This server row moved or changed in the master workbook. Re-import it before editing.")

        header_indexes = {heading: index + 1 for index, heading in enumerate(headers)}
        inventory_indexes = {heading: index + 1 for index, heading in enumerate(inventory_headers)}
        updated_values = {}
        for heading, text in updates.items():
            if heading in inventory_indexes:
                target_sheet = inventory
                column = inventory_indexes[heading]
            else:
                target_sheet = sheet
                column = header_indexes.get(heading)
                if column is None:
                    raise ImportValidationError(f"The heading '{heading}' is not present in the master worksheet.")
            cell = target_sheet.cell(row=row, column=column)
            if cell.data_type == "f":
                raise ImportValidationError(f"{heading} is a formula cell and cannot be changed from this page.")
            value = _cell_value_from_text(text, cell.value, heading) if isinstance(text, str) else text
            if isinstance(value, str) and value.lstrip().startswith("="):
                raise ImportValidationError(f"Formula-like values are not allowed in {heading}. Enter plain text only.")
            cell.value = value
            if isinstance(value, datetime) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd"
            updated_values[heading] = value

        if getattr(workbook, "calculation", None):
            workbook.calculation.calcMode = "auto"
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), updated_values
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The master workbook row could not be updated safely.") from exc
    finally:
        workbook.close()


def _atomic_replace(path, content):
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", dir=path.parent, prefix=f".{path.stem}-", suffix=".xlsx", delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
            temporary.write(content)
            temporary.flush()
            os.fsync(temporary.fileno())
        os.replace(temporary_path, path)
    except PermissionError as exc:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink(missing_ok=True)
        raise ImportValidationError(
            "Windows could not replace the original workbook. Close it in Excel and save the web edit again."
        ) from exc
    except Exception:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink(missing_ok=True)
        raise


def _value(record_data, heading):
    value = record_data.get(heading)
    return "" if value is None else str(value)


def _clean_ip(value):
    value = (value or "").strip()
    return "" if value.casefold() in {"nan", "not found", "n/a", "na", "none"} else value


def _update_server_fields(record, data, batch):
    hostname = (_value(data, "Host Name") or _value(data, "Hostname")).strip()
    if not hostname:
        raise ImportValidationError("Host Name cannot be blank because it identifies the server.")
    normalized = hostname.casefold()
    server = record.server
    if server and hostname != record.hostname and batch.snapshot_date < (server.last_seen_on or batch.snapshot_date):
        raise ImportValidationError("Host Name can only be changed in the latest snapshot for this server.")
    if server and hostname != record.hostname and ImportRecord.objects.filter(batch=batch, server=server).exclude(pk=record.pk).exists():
        raise ImportValidationError("This server has multiple rows in the snapshot. Edit their Host Name values together.")
    if Server.objects.filter(normalized_hostname=normalized).exclude(pk=server.pk if server else None).exists():
        raise ImportValidationError("That Host Name already belongs to another server. Merge the records manually first.")
    if server is None:
        server = Server(
            hostname=hostname,
            normalized_hostname=normalized,
            first_seen_on=batch.snapshot_date,
            last_seen_on=batch.snapshot_date,
        )
        server.save()
        record.server = server
    else:
        server.hostname = hostname
        server.normalized_hostname = normalized

    record.hostname = hostname
    record.ip_address = _clean_ip(_value(data, "IP Address"))
    record.operating_system = _value(data, "OPERATING SYSTEM").strip()
    record.environment = _value(data, "Environment").strip()

    if server.last_seen_on is None or batch.snapshot_date >= server.last_seen_on:
        peers = list(ImportRecord.objects.filter(batch=batch, server=server).exclude(pk=record.pk))
        peer_values = [peer.source_data for peer in peers] + [data]
        ips = {_clean_ip(_value(item, "IP Address")) for item in peer_values if _clean_ip(_value(item, "IP Address"))}
        systems = {_value(item, "OPERATING SYSTEM").strip() for item in peer_values if _value(item, "OPERATING SYSTEM").strip()}
        environments = {_value(item, "Environment").strip() for item in peer_values if _value(item, "Environment").strip()}
        server.current_ip = next(iter(ips)) if len(ips) == 1 else ""
        server.ip_ambiguous = len(ips) > 1
        server.operating_system = next(iter(systems)) if len(systems) == 1 else ""
        server.environment = next(iter(environments)) if len(environments) == 1 else ""
        if server.first_seen_on is None or batch.snapshot_date < server.first_seen_on:
            server.first_seen_on = batch.snapshot_date
        if server.last_seen_on is None or batch.snapshot_date > server.last_seen_on:
            server.last_seen_on = batch.snapshot_date
    server.save()
    record.save(update_fields=[
        "server", "hostname", "ip_address", "operating_system", "environment", "source_data",
        "raw_status", "mapped_status",
    ])
    return server


def _record_status_change(record, batch, actor, previous_status, new_status, raw_status):
    today = timezone.localdate()
    state = ComplianceState.objects.filter(server=record.server, parameter=batch.parameter).first()
    if state and previous_status == new_status:
        if state.raw_status != raw_status:
            state.raw_status = raw_status
            state.latest_import = batch
            state.observed_on = today
            state.save(update_fields=["raw_status", "latest_import", "observed_on", "updated_at"])
        return
    was_compliant = bool(state and state.status == ComplianceStatus.COMPLIANT)
    recurring = bool(
        was_compliant
        and new_status == ComplianceStatus.NON_COMPLIANT
        and StatusEvent.objects.filter(
            server=record.server,
            parameter=batch.parameter,
            new_status=ComplianceStatus.NON_COMPLIANT,
        ).exists()
    )
    event_previous_status = previous_status
    if state is None:
        event_previous_status = ""
        state = ComplianceState(
            server=record.server,
            parameter=batch.parameter,
            status=new_status,
            observed_on=today,
        )
    state.status = new_status
    state.raw_status = raw_status
    state.latest_import = batch
    state.observed_on = today
    if recurring:
        state.recurrence_count += 1
    state.save()
    StatusEvent.objects.create(
        server=record.server,
        parameter=batch.parameter,
        import_batch=batch,
        previous_status=event_previous_status,
        new_status=new_status,
        kind=StatusEvent.Kind.RECURRENCE if recurring else StatusEvent.Kind.MANUAL,
        event_date=today,
        reason=f"Excel row {record.source_row} edited by {actor}.",
    )


def edit_source_record(record_id, updates, actor):
    """Update one row in the stored workbook, optional original, and normalized data."""
    record = ImportRecord.objects.select_related("batch", "batch__parameter", "server").get(pk=record_id)
    batch = record.batch
    latest_batch = ImportBatch.objects.filter(
        parameter=batch.parameter,
        is_current_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    if latest_batch is None or latest_batch.pk != batch.pk:
        raise ImportValidationError("Only rows in the latest snapshot can be edited. Re-import a newer workbook first.")

    master_config = MasterWorkbook.objects.first() if batch.is_master_source else None
    target = None
    target_original = None
    if master_config and master_config.path:
        target = Path(master_config.path)
        try:
            source_content = target.read_bytes()
        except OSError as exc:
            raise ImportValidationError("The configured master workbook is no longer available on this computer.") from exc
        if sha256(source_content).hexdigest() != master_config.content_sha256:
            raise ImportValidationError("The master workbook changed outside the web app. Re-import the latest file before editing.")
        target_original = source_content
    else:
        source_content = _read_uploaded_copy(batch)
    if not batch.is_master_source and batch.sync_target_path:
        path = validate_sync_target_path(
            batch.sync_target_path,
            expected_hash=batch.working_sha256 or batch.file_sha256,
            source_sheet=batch.source_sheet,
            source_columns=batch.source_columns,
        )
        target = Path(path)
        target_original = target.read_bytes()

    now = timezone.localtime().replace(tzinfo=None, second=0, microsecond=0)
    values = dict(record.source_data)
    actual_updates = dict(updates)
    for heading in batch.source_columns:
        key = "".join(character for character in heading.casefold() if character.isalnum())
        if key in {"lastupdated", "lastupdatedon", "lastmodified", "checkdate"}:
            actual_updates[heading] = now
            values[heading] = now.isoformat(sep=" ")
        elif key in {"changesmadeby", "changedby", "updatedby"}:
            actual_updates[heading] = actor
            values[heading] = actor
    for heading, value in actual_updates.items():
        if isinstance(value, (datetime, date)):
            values[heading] = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
        else:
            values[heading] = value

    if batch.status_column in actual_updates and record.server_id:
        peer_data = batch.records.filter(server_id=record.server_id).exclude(pk=record.pk).values_list("source_data", flat=True)
        resulting_statuses = {
            "" if item.get(batch.status_column) is None else str(item.get(batch.status_column))
            for item in peer_data
        }
        resulting_statuses.add("" if values.get(batch.status_column) is None else str(values.get(batch.status_column)))
        if len(resulting_statuses) > 1:
            raise ImportValidationError(
                "This server has multiple rows in the snapshot. Keep their status values consistent when editing."
            )

    if batch.is_master_source and target:
        updated_content, written_values = _rewrite_master_workbook(source_content, record, actual_updates)
    else:
        updated_content, written_values = _rewrite_workbook(source_content, record, actual_updates)
    new_hash = sha256(updated_content).hexdigest()
    working_file = batch.working_file or batch.source_file
    storage = working_file.storage
    old_name = working_file.name
    original_name = batch.source_file.name
    new_name = storage.save(
        f"imports/edited/{uuid.uuid4().hex}.xlsx",
        ContentFile(updated_content),
    )
    external_updated = False
    old_working_names = set()
    try:
        with transaction.atomic():
            record = ImportRecord.objects.select_for_update().select_related("server", "batch", "batch__parameter").get(pk=record_id)
            batch = ImportBatch.objects.select_for_update().get(pk=record.batch_id)
            if batch.is_master_source and master_config:
                current_master = MasterWorkbook.objects.select_for_update().get(pk=master_config.pk)
                if current_master.path != str(target) or current_master.content_sha256 != sha256(source_content).hexdigest():
                    raise ImportValidationError("The master workbook settings changed while this row was being edited. Reload and try again.")
                master_config = current_master
            if not batch.is_master_source and (batch.working_sha256 or batch.file_sha256) != sha256(source_content).hexdigest():
                raise ImportValidationError("This row changed while it was being edited. Reload the page and try again.")
            if target:
                expected_hash = master_config.content_sha256 if batch.is_master_source and master_config else (batch.working_sha256 or batch.file_sha256)
                if sha256(target.read_bytes()).hexdigest() != expected_hash:
                    raise ImportValidationError(
                        "The workbook changed outside the web app. Re-import the latest workbook before editing again."
                    )
                if batch.is_master_source:
                    backup = target.with_name(f"{target.stem}.before_web_edits{target.suffix}")
                    if not backup.exists():
                        with backup.open("xb") as handle:
                            handle.write(target_original)
                            handle.flush()
                            os.fsync(handle.fileno())
                _atomic_replace(target, updated_content)
                external_updated = True

            previous_status = record.mapped_status
            previous_data = dict(record.source_data)
            record.source_data = values
            for heading, value in written_values.items():
                if isinstance(value, datetime):
                    record.source_data[heading] = value.isoformat(sep=" ")
                elif isinstance(value, date):
                    record.source_data[heading] = value.isoformat()
                else:
                    record.source_data[heading] = value
            status_value = values.get(batch.status_column, "")
            raw_status = "" if status_value is None else str(status_value)
            new_status = STANDARD_STATUS_MAP.get(raw_status.strip(), ComplianceStatus.UNMAPPED)
            record.raw_status = raw_status
            record.mapped_status = new_status
            server = _update_server_fields(record, record.source_data, batch)
            record.server = server
            record.save(update_fields=["source_data", "raw_status", "mapped_status", "server", "hostname", "ip_address", "operating_system", "environment"])
            _record_status_change(record, batch, actor, previous_status, new_status, raw_status)

            if batch.is_master_source and master_config:
                master_config = MasterWorkbook.objects.select_for_update().get(pk=master_config.pk)
                master_config.content_sha256 = new_hash
                master_config.save(update_fields=["content_sha256", "updated_at"])
                current_batches = ImportBatch.objects.select_for_update().filter(
                    is_master_source=True, is_current_source=True
                )
                old_working_names = {
                    item for item in current_batches.values_list("working_file", flat=True) if item
                }
                current_batches.update(working_file=new_name, working_sha256=new_hash)
                batch.working_file.name = new_name
                batch.working_sha256 = new_hash
            else:
                batch.working_file.name = new_name
                batch.working_sha256 = new_hash
                batch.save(update_fields=["working_file", "working_sha256"])
            changed_values = {
                heading: {"from": previous_data.get(heading), "to": values.get(heading)}
                for heading in actual_updates if heading in previous_data
            }
            changed_summary = json.dumps(changed_values, ensure_ascii=False, default=str)[:8000]
            AuditLog.objects.create(
                action="excel_row_edited",
                detail=f"Excel row {record.source_row} edited by {actor}. Changes: {changed_summary or 'No source values changed.'}",
                import_batch=batch,
                server=server,
            )
    except Exception:
        storage.delete(new_name)
        if external_updated and target and target_original is not None:
            _atomic_replace(target, target_original)
        raise
    if batch.is_master_source:
        if target and master_config:
            # Reconcile every dated status column after an edit. The batch currently
            # being edited may represent the newest month while an older month cell
            # in the same worksheet also feeds historical comparison charts.
            from website.services.master_workbook import import_master_workbook

            import_master_workbook(
                updated_content,
                original_filename=master_config.original_filename,
                target_path=str(target),
            )
        for old_file in old_working_names:
            if old_file == new_name:
                continue
            if not ImportBatch.objects.filter(source_file=old_file).exists() and not ImportBatch.objects.filter(working_file=old_file).exists():
                storage.delete(old_file)
    if old_name and old_name != original_name:
        storage.delete(old_name)
    return record
