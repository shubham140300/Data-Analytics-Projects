"""Owner-controlled optional column changes for an imported Excel snapshot."""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import uuid

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from website.models import AuditLog, ImportBatch, ImportRecord, MasterWorkbook
from website.services.errors import ImportValidationError
from website.services.excel_row_edit import _atomic_replace, _read_uploaded_copy


LOCKED_HEADINGS = {
    "hostname", "ipaddress", "serverid", "parameter", "status", "operatingsystem", "environment",
    "lastupdated", "lastupdatedon", "lastmodified", "changesmadeby", "changedby", "updatedby", "checkdate",
}


def _key(value):
    return "".join(character for character in value.casefold() if character.isalnum())


def _heading_indexes(headers):
    return {heading: index + 1 for index, heading in enumerate(headers)}


def _master_locked(heading):
    return _key(heading) in LOCKED_HEADINGS or "status" in _key(heading)


def _save_master_workbook_columns(batch_id, operation, existing, new):
    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    config = MasterWorkbook.objects.select_for_update().first()
    if not config or not config.path:
        raise ImportValidationError("The master workbook write-back path is not configured.")
    target = Path(config.path)
    output = BytesIO()
    try:
        source_content = target.read_bytes()
    except OSError as exc:
        raise ImportValidationError("The configured master workbook is no longer available.") from exc
    if sha256(source_content).hexdigest() != config.content_sha256:
        raise ImportValidationError("The master workbook changed outside the app. Upload the latest workbook before changing columns.")

    headers = list(batch.source_columns)
    if not headers or len({item.casefold() for item in headers}) != len(headers):
        raise ImportValidationError("The workbook has duplicate or missing headings. Upload the latest master workbook first.")
    if operation in {"rename", "delete"}:
        if existing not in headers:
            raise ImportValidationError("That heading is no longer in the master workbook. Reload and try again.")
        if _master_locked(existing):
            raise ImportValidationError("Server identity, parameter, status, and audit columns cannot be changed here.")
    if operation in {"add", "rename"}:
        collisions = [item for item in headers if item != existing]
        if not new or any(item.casefold() == new.casefold() for item in collisions):
            raise ImportValidationError("Column names must be non-empty and unique, ignoring letter case.")
        if len(new) > 255 or new.lstrip().startswith("=") or any(ord(char) < 32 for char in new):
            raise ImportValidationError("Column names must be plain text, at most 255 characters, with no control characters.")
    if operation not in {"add", "rename", "delete"}:
        raise ImportValidationError("Choose a supported column operation.")

    try:
        workbook = load_workbook(BytesIO(source_content), data_only=False)
    except Exception as exc:
        raise ImportValidationError("The current master workbook could not be opened.") from exc
    try:
        if batch.source_sheet not in workbook.sheetnames:
            raise ImportValidationError("The master worksheet for this parameter is missing.")
        sheet = workbook[batch.source_sheet]
        current_headers = ["" if cell.value is None else str(cell.value).strip() for cell in sheet[1]]
        while current_headers and not current_headers[-1]:
            current_headers.pop()
        if current_headers != headers:
            raise ImportValidationError("The master headings changed. Upload the latest workbook before editing columns.")

        column_index = _heading_indexes(headers).get(existing) if existing else None
        if operation == "add":
            headers.append(new)
            sheet.cell(row=1, column=len(headers)).value = new
        elif operation == "rename":
            sheet.cell(row=1, column=column_index).value = new
            headers[headers.index(existing)] = new
        else:
            target_letter = get_column_letter(column_index)
            if any(sheet.tables.values()):
                raise ImportValidationError("This sheet is part of an Excel table. Remove the column in Excel and upload the updated master.")
            if sheet.merged_cells.ranges or sheet._charts or sheet.data_validations.dataValidation:
                raise ImportValidationError("This sheet has layout features that could be damaged. Remove the column in Excel and upload the updated master.")
            reference = re.compile(
                rf"(?:'{re.escape(sheet.title)}'|{re.escape(sheet.title)})!\$?{re.escape(target_letter)}\$?\d+",
                re.IGNORECASE,
            )
            for formula_sheet in workbook.worksheets:
                for row in formula_sheet.iter_rows():
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        if reference.search(str(cell.value)):
                            raise ImportValidationError("A workbook formula uses this column. Remove it in Excel and upload the updated master.")
                        if formula_sheet.title == sheet.title and re.search(
                            rf"(?<![A-Z0-9_])\$?{re.escape(target_letter)}\$?\d+",
                            str(cell.value), re.IGNORECASE,
                        ):
                            raise ImportValidationError("A worksheet formula uses this column. Remove it in Excel and upload the updated master.")
            sheet.delete_cols(column_index, 1)
            headers.remove(existing)
        if getattr(workbook, "calculation", None):
            workbook.calculation.calcMode = "auto"
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(output)
        updated_content = output.getvalue()
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The master workbook column could not be updated safely.") from exc
    finally:
        workbook.close()

    new_hash = sha256(updated_content).hexdigest()

    related_batches = list(ImportBatch.objects.select_for_update().filter(
        is_current_source=True,
        is_master_source=True,
        source_sheet=batch.source_sheet,
    ))
    records = list(ImportRecord.objects.select_for_update().filter(
        batch__in=related_batches
    ).only("id", "source_data"))
    for record in records:
        values = dict(record.source_data)
        if operation == "add":
            values[new] = None
        elif operation == "rename":
            values[new] = values.pop(existing, None)
        else:
            values.pop(existing, None)
        record.source_data = values

    stored_file = batch.working_file or batch.source_file
    storage = stored_file.storage
    old_names = {item for item in related_batches for item in (item.working_file.name,) if item.working_file}
    new_file_name = storage.save(f"imports/edited/{uuid.uuid4().hex}.xlsx", ContentFile(updated_content))
    target_original = source_content
    target_updated = False
    try:
        if sha256(target.read_bytes()).hexdigest() != config.content_sha256:
            raise ImportValidationError("The master workbook changed while its columns were being edited. Reload and try again.")
        backup = target.with_name(f"{target.stem}.before_web_edits{target.suffix}")
        if not backup.exists():
            with backup.open("xb") as handle:
                handle.write(target_original)
                handle.flush()
        _atomic_replace(target, updated_content)
        target_updated = True
        with transaction.atomic():
            locked_config = MasterWorkbook.objects.select_for_update().get(pk=config.pk)
            if locked_config.content_sha256 != sha256(source_content).hexdigest():
                raise ImportValidationError("The master workbook settings changed while columns were being edited.")
            locked_config.content_sha256 = new_hash
            locked_config.save(update_fields=["content_sha256", "updated_at"])
            ImportRecord.objects.bulk_update(records, ["source_data"], batch_size=500)
            for item in related_batches:
                item.source_columns = headers
                item.working_file.name = new_file_name
                item.working_sha256 = new_hash
            ImportBatch.objects.bulk_update(
                related_batches,
                ["source_columns", "working_file", "working_sha256"],
                batch_size=100,
            )
            AuditLog.objects.create(
                action=f"master_column_{operation}",
                detail=f"Owner applied '{operation}' to {batch.source_sheet}: {existing or new}"
                + (f" → {new}" if operation == "rename" else ""),
                import_batch=batch,
            )
    except Exception:
        storage.delete(new_file_name)
        if target_updated:
            _atomic_replace(target, target_original)
        raise
    for old_name in old_names:
        if old_name != new_file_name and not ImportBatch.objects.filter(source_file=old_name).exists() and not ImportBatch.objects.filter(working_file=old_name).exists():
            storage.delete(old_name)
    return ImportBatch.objects.get(pk=batch_id)


def _save_workbook(batch_id, operation, existing, new):
    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    expected_hash = batch.working_sha256 or batch.file_sha256
    source_content = _read_uploaded_copy(batch)
    if sha256(source_content).hexdigest() != expected_hash:
        raise ImportValidationError("The app's workbook copy changed outside the web app. Re-import it before changing columns.")

    headers = list(batch.source_columns)
    if not headers or len(set(header.casefold() for header in headers)) != len(headers):
        raise ImportValidationError("This workbook has duplicate or missing headings. Re-import it before changing columns.")
    if operation in {"rename", "delete"}:
        if existing not in headers:
            raise ImportValidationError("That column is no longer in this workbook. Reload the page and try again.")
        if existing == batch.status_column or _key(existing) in LOCKED_HEADINGS:
            raise ImportValidationError("Core server identity, status, and audit columns cannot be changed.")
    if operation in {"add", "rename"}:
        collision_headers = [header for header in headers if header != existing]
        if not new or any(header.casefold() == new.casefold() for header in collision_headers):
            raise ImportValidationError("Column names must be non-empty and unique, ignoring letter case.")
        if len(new) > 255:
            raise ImportValidationError("Column names can have up to 255 characters.")
        if new.lstrip().startswith("=") or any(ord(character) < 32 for character in new):
            raise ImportValidationError("Column names cannot contain formulas or control characters.")
    if operation not in {"add", "rename", "delete"}:
        raise ImportValidationError("Choose a supported column operation.")

    try:
        workbook = load_workbook(BytesIO(source_content), data_only=False)
    except Exception as exc:
        raise ImportValidationError("The saved .xlsx workbook could not be opened.") from exc
    try:
        if batch.source_sheet not in workbook.sheetnames:
            raise ImportValidationError("The imported worksheet is missing from the saved workbook.")
        sheet = workbook[batch.source_sheet]
        current_headers = ["" if cell.value is None else str(cell.value) for cell in sheet[1]]
        while current_headers and not current_headers[-1].strip():
            current_headers.pop()
        if current_headers != headers:
            raise ImportValidationError("Workbook headings no longer match the imported snapshot. Re-import the latest workbook.")

        records = list(ImportRecord.objects.filter(batch=batch).only("id", "source_data"))
        if operation == "add":
            headers.append(new)
            sheet.cell(row=1, column=len(headers)).value = new
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row=row_number, column=len(headers)).value = None
            for record in records:
                record.source_data = {**record.source_data, new: None}
        elif operation == "rename":
            column_index = _heading_indexes(headers)[existing]
            sheet.cell(row=1, column=column_index).value = new
            headers[headers.index(existing)] = new
            for record in records:
                values = dict(record.source_data)
                values[new] = values.pop(existing, None)
                record.source_data = values
        else:
            column_index = _heading_indexes(headers)[existing]
            has_formulas = any(
                cell.data_type == "f"
                for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column)
                for cell in row
            )
            if has_formulas:
                raise ImportValidationError("This workbook contains formulas. Remove the column in Excel and re-import to preserve formulas safely.")
            if sheet.tables or sheet.merged_cells.ranges or sheet._charts or sheet.data_validations.dataValidation:
                raise ImportValidationError(
                    "This workbook has tables, merged cells, charts, or data validation. Remove the column in Excel and re-import to preserve its layout."
                )
            sheet.delete_cols(column_index, 1)
            headers.remove(existing)
            for record in records:
                values = dict(record.source_data)
                values.pop(existing, None)
                record.source_data = values

        output = BytesIO()
        workbook.save(output)
        updated_content = output.getvalue()
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The workbook columns could not be updated safely.") from exc
    finally:
        workbook.close()

    new_hash = sha256(updated_content).hexdigest()
    target = None
    target_original = None
    target_updated = False
    if batch.sync_target_path:
        from website.services.excel_row_edit import validate_sync_target_path

        target_path = validate_sync_target_path(
            batch.sync_target_path,
            expected_hash=expected_hash,
            source_sheet=batch.source_sheet,
            source_columns=batch.source_columns,
        )
        target = Path(target_path)
        target_original = target.read_bytes()
        if sha256(target_original).hexdigest() != expected_hash:
            raise ImportValidationError("The original workbook changed outside the app. Re-import before changing columns.")

    stored_file = batch.working_file or batch.source_file
    storage = stored_file.storage
    old_name = stored_file.name
    original_name = batch.source_file.name
    new_file_name = storage.save(f"imports/edited/{uuid.uuid4().hex}.xlsx", ContentFile(updated_content))
    try:
        with transaction.atomic():
            batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
            if (batch.working_sha256 or batch.file_sha256) != expected_hash:
                raise ImportValidationError("The workbook changed while columns were being updated. Reload and try again.")
            if target:
                _atomic_replace(target, updated_content)
                target_updated = True
            ImportRecord.objects.bulk_update(records, ["source_data"], batch_size=500)
            batch.source_columns = headers
            batch.working_file.name = new_file_name
            batch.working_sha256 = new_hash
            batch.save(update_fields=["source_columns", "working_file", "working_sha256"])
            AuditLog.objects.create(
                action=f"excel_column_{operation}",
                detail=f"Column operation '{operation}' applied to workbook headings: {existing or new}"
                + (f" → {new}" if operation == "rename" else ""),
                import_batch=batch,
            )
    except Exception:
        storage.delete(new_file_name)
        if target_updated and target and target_original is not None:
            _atomic_replace(target, target_original)
        raise
    if old_name and old_name != original_name:
        storage.delete(old_name)
    return batch


def change_workbook_columns(batch_id, *, operation, existing_heading="", new_heading=""):
    """Apply a schema change to the stored workbook, import rows, and optional original file."""
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
        if batch.is_master_source:
            return _save_master_workbook_columns(
                batch_id=batch_id,
                operation=operation,
                existing=existing_heading,
                new=new_heading,
            )
        return _save_workbook(
            batch_id=batch_id,
            operation=operation,
            existing=existing_heading,
            new=new_heading,
        )
