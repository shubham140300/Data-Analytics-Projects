from io import BytesIO
from datetime import timedelta
from hashlib import sha256
import ipaddress
import json
from pathlib import Path
from collections import defaultdict

from django.contrib import messages
from django.conf import settings
from django.contrib.auth.hashers import check_password, make_password
from django.core.paginator import Paginator
from django.db.models import Count, Max, Q
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden, HttpResponseNotAllowed
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.text import slugify
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from .forms import (
    ColumnChangeForm,
    ExcelImportForm,
    MasterWorkbookImportForm,
    ManualStateChangeForm,
    MasterSheetRowEditForm,
    SourceRecordEditForm,
    SyncTargetForm,
    WorkspacePasswordChangeForm,
    WorkspaceUserGrantForm,
)
from .access import allowed_parameter_ids, has_parameter_permission
from .models import (
    AuditLog,
    ComplianceState,
    ComplianceStatus,
    ImportBatch,
    ImportRecord,
    MasterWorkbook,
    Parameter,
    Server,
    StatusEvent,
    WorkspaceUser,
    WorkspaceParameterPermission,
)
from .services.errors import ImportValidationError
from .services.excel_row_edit import edit_source_record, validate_sync_target_path
from .services.excel_import import preview_workbook
from .services.import_workflow import import_snapshot
from .services.excel_columns import change_workbook_columns, LOCKED_HEADINGS
from .services.master_workbook import import_master_workbook, parameter_sheet_name, read_master_parameter_sheet
from .services.master_sheet_edit import edit_master_sheet_row
from .services.history_admin import clear_compliance_workspace_data, delete_import_batch, delete_manual_status_event


PAGE_SIZES = {"25": 25, "50": 50, "100": 100}
QUALITY_LABELS = {
    "missing_host": "Missing Host Name",
    "missing_ip": "Missing IP Address",
    "missing_os": "Missing operating system",
    "missing_environment": "Missing environment",
    "invalid_ip": "Invalid IP address format",
    "unmapped_status": "Unmapped status",
    "conflicting_status": "Conflicting statuses",
    "duplicate_row": "Duplicate source row",
}


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _record_signature(record):
    serialized = json.dumps(record.source_data or {}, ensure_ascii=False, sort_keys=True, default=str)
    return sha256(serialized.encode("utf-8")).hexdigest()


def _is_valid_ip(value):
    try:
        ipaddress.ip_interface(str(value).strip())
        return True
    except ValueError:
        return False


def _dashboard_data_quality(allowed_ids=None, writable_ids=None):
    latest_batches = []
    parameters = Parameter.objects.all()
    master = MasterWorkbook.objects.first()
    if master and master.parameter_names:
        parameters = parameters.filter(name__in=master.parameter_names)
    if allowed_ids is not None:
        parameters = parameters.filter(pk__in=allowed_ids)
    for parameter in parameters:
        batch = ImportBatch.objects.select_related("parameter").filter(
            parameter=parameter,
            is_current_source=True,
            result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
        ).order_by("-snapshot_date", "-imported_at").first()
        if batch:
            latest_batches.append(batch)
    latest_batches.sort(key=lambda batch: (batch.snapshot_date, batch.imported_at), reverse=True)

    counts = {key: 0 for key in QUALITY_LABELS}
    examples = []
    issue_row_count = 0
    missing_data_row_count = 0
    other_quality_row_count = 0
    missing_keys = {"missing_host", "missing_ip", "missing_os", "missing_environment"}

    for batch in latest_batches:
        records = ImportRecord.objects.filter(batch=batch).only(
            "id", "source_row", "server_id", "hostname", "ip_address", "operating_system",
            "environment", "raw_status", "mapped_status", "source_data",
        ).order_by("source_row")
        status_variants = defaultdict(set)
        seen_signatures = set()
        duplicate_signatures = set()
        for record in records.iterator(chunk_size=1000):
            hostname_key = record.hostname.strip().casefold()
            if hostname_key:
                status_variants[hostname_key].add((record.mapped_status, record.raw_status))
            signature = _record_signature(record)
            if signature in seen_signatures:
                duplicate_signatures.add(signature)
            seen_signatures.add(signature)
        conflicting_hostnames = {
            hostname for hostname, variants in status_variants.items() if len(variants) > 1
        }

        seen_signatures = set()
        for record in records.iterator(chunk_size=1000):
            signature = _record_signature(record)
            is_duplicate = signature in seen_signatures
            seen_signatures.add(signature)
            hostname_key = record.hostname.strip().casefold()
            issue_keys = []
            if _blank(record.hostname):
                issue_keys.append("missing_host")
            if _blank(record.ip_address):
                issue_keys.append("missing_ip")
            if not batch.is_master_source and _blank(record.operating_system):
                issue_keys.append("missing_os")
            if not batch.is_master_source and _blank(record.environment):
                issue_keys.append("missing_environment")
            if not _blank(record.ip_address) and not _is_valid_ip(record.ip_address):
                issue_keys.append("invalid_ip")
            if record.mapped_status == ComplianceStatus.UNMAPPED:
                issue_keys.append("unmapped_status")
            if hostname_key and hostname_key in conflicting_hostnames:
                issue_keys.append("conflicting_status")
            if is_duplicate and signature in duplicate_signatures:
                issue_keys.append("duplicate_row")
            if not issue_keys:
                continue

            issue_row_count += 1
            for key in issue_keys:
                counts[key] += 1
            if missing_keys.intersection(issue_keys):
                missing_data_row_count += 1
            if set(issue_keys) - missing_keys:
                other_quality_row_count += 1
            if len(examples) < 12:
                examples.append({
                    "id": record.id,
                    "source_row": record.source_row,
                    "batch": batch,
                    "parameter": batch.parameter.name,
                    "parameter_id": batch.parameter_id,
                    "can_edit": writable_ids is None or batch.parameter_id in writable_ids,
                    "hostname": record.hostname,
                    "ip_address": record.ip_address,
                    "issues": [QUALITY_LABELS[key] for key in issue_keys],
                })

    return {
        "quality_categories": [
            {"key": key, "label": label, "count": counts[key]}
            for key, label in QUALITY_LABELS.items()
        ],
        "quality_issue_count": issue_row_count,
        "quality_snapshot_count": len(latest_batches),
        "missing_data_row_count": missing_data_row_count,
        "other_quality_row_count": other_quality_row_count,
        "quality_examples": examples,
    }


def login(request):
    if request.workspace_user_id:
        return redirect("dashboard")
    error = ""
    if request.method == "POST":
        username = request.POST.get("username", "").strip().casefold()
        password = request.POST.get("password", "")
        grant = WorkspaceUser.objects.filter(username__iexact=username, is_active=True).first()
        is_owner_account = bool(grant and grant.user_id == settings.WORKSPACE_OWNER_ID)
        if (
            grant
            and (is_owner_account or grant.can_read)
            and check_password(password, grant.password_hash)
            and (not is_owner_account or grant.username == settings.WORKSPACE_OWNER_USERNAME)
        ):
            request.session.cycle_key()
            request.session["workspace_user_id"] = grant.user_id
            request.session["workspace_auth_method"] = "password"
            request.session["workspace_must_change_password"] = grant.must_change_password
            request.session.set_expiry(
                settings.WORKSPACE_OWNER_SESSION_AGE if is_owner_account else 8 * 60 * 60
            )
            return redirect(_safe_next(request))
        else:
            if username == settings.WORKSPACE_OWNER_USERNAME and not grant:
                error = "Set up the owner account locally first with: python manage.py setup_workspace_owner"
            else:
                error = "Invalid username or password, or this account is not active. Ask the owner to check your access."
    return render(request, "website/login.html", {"error": error, "next_url": request.GET.get("next", "")})


def _safe_next(request):
    target = request.POST.get("next", "") or request.GET.get("next", "")
    if url_has_allowed_host_and_scheme(target, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return target
    return reverse("dashboard")


def logout(request):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    request.session.flush()
    return redirect("login")


def password_change(request):
    grant = request.workspace_user
    form = WorkspacePasswordChangeForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        if not check_password(form.cleaned_data["current_password"], grant.password_hash):
            form.add_error("current_password", "The current password is incorrect.")
        else:
            grant.password_hash = make_password(form.cleaned_data["new_password"])
            grant.must_change_password = False
            grant.save(update_fields=["password_hash", "must_change_password", "updated_at"])
            request.session.pop("workspace_must_change_password", None)
            messages.success(request, "Your password was changed.")
            return redirect("dashboard")
    return render(request, "website/password_change.html", {"form": form})


def access_control(request):
    selected_id = request.GET.get("user", "").strip()
    instance = WorkspaceUser.objects.filter(user_id=selected_id).exclude(
        user_id=settings.WORKSPACE_OWNER_ID
    ).first() if selected_id else None
    if request.method == "POST" and request.POST.get("action") == "deactivate":
        grant = get_object_or_404(
            WorkspaceUser.objects.exclude(user_id=settings.WORKSPACE_OWNER_ID),
            user_id=request.POST.get("user_id", ""),
        )
        grant.is_active = False
        grant.save(update_fields=["is_active", "updated_at"])
        AuditLog.objects.create(
            action="workspace_access_revoked",
            detail=f"Access disabled for user ID {grant.user_id} by owner {request.workspace_user_id}.",
        )
        messages.success(request, f"Access for user ID {grant.user_id} was disabled.")
        return redirect("access_control")

    master = MasterWorkbook.objects.first()
    parameters = Parameter.objects.filter(name__in=master.parameter_names) if master and master.parameter_names else Parameter.objects.all()
    form = WorkspaceUserGrantForm(request.POST or None, instance=instance, parameters=parameters)
    if request.method == "POST" and form.is_valid():
        grant = form.save(granted_by=request.workspace_user_id)
        AuditLog.objects.create(
            action="workspace_access_granted" if not instance else "workspace_access_updated",
            detail=(
                f"Per-parameter access for user ID {grant.user_id} updated; active={grant.is_active}."
            ),
        )
        messages.success(request, f"Access settings saved for user ID {grant.user_id}.")
        return redirect("access_control")

    grants = WorkspaceUser.objects.exclude(user_id=settings.WORKSPACE_OWNER_ID).prefetch_related("parameter_permissions__parameter")
    return render(request, "website/access_control.html", {
        "form": form,
        "grants": grants,
        "editing_grant": instance,
        "parameters": parameters,
        "default_password": settings.WORKSPACE_DEFAULT_PASSWORD,
    })


def column_settings(request, batch_id):
    batch = get_object_or_404(ImportBatch.objects.select_related("parameter"), pk=batch_id)
    latest_batch = ImportBatch.objects.filter(
        parameter=batch.parameter,
        is_current_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    if latest_batch is None or latest_batch.pk != batch.pk:
        messages.error(request, "Column settings are available for the latest snapshot of each parameter.")
        return redirect("import_detail", batch_id=batch.pk)
    locked_columns = [
        heading for heading in batch.source_columns
        if (
            (
                "status" in "".join(c for c in heading.casefold() if c.isalnum())
                or "".join(c for c in heading.casefold() if c.isalnum()) in LOCKED_HEADINGS
            )
            if batch.is_master_source
            else (
                heading == batch.status_column
                or "".join(c for c in heading.casefold() if c.isalnum()) in LOCKED_HEADINGS
            )
        )
    ]
    optional_columns = [heading for heading in batch.source_columns if heading not in locked_columns]
    form = ColumnChangeForm(request.POST or None)
    form.fields["existing_heading"].choices = [("", "Choose a column")] + [(heading, heading) for heading in optional_columns]
    if request.method == "POST" and form.is_valid():
        try:
            batch = change_workbook_columns(
                batch.pk,
                operation=form.cleaned_data["operation"],
                existing_heading=form.cleaned_data["existing_heading"],
                new_heading=form.cleaned_data["new_heading"],
            )
        except ImportValidationError as exc:
            form.add_error(None, str(exc))
        except OSError:
            form.add_error(None, "The workbook could not be saved. Close it in Excel and try again.")
        else:
            messages.success(request, "The column change was saved to the app workbook and configured original workbook.")
            return redirect("column_settings", batch_id=batch.pk)
    return render(request, "website/column_settings.html", {
        "batch": batch,
        "form": form,
        "locked_columns": locked_columns,
        "optional_columns": optional_columns,
    })


def _paginate(request, items):
    requested_size = request.GET.get("page_size", "25")
    page_size = PAGE_SIZES.get(requested_size, 25)
    return Paginator(items, page_size).get_page(request.GET.get("page")), page_size


def dashboard(request):
    readable_ids = allowed_parameter_ids(request, "read")
    master = MasterWorkbook.objects.first()
    states = ComplianceState.objects.filter(latest_import__is_current_source=True)
    if readable_ids is not None:
        states = states.filter(parameter_id__in=readable_ids)
    status_counts = dict(states.values_list("status").annotate(count=Count("id")).values_list("status", "count"))
    compliant_count = status_counts.get(ComplianceStatus.COMPLIANT, 0)
    non_compliant_count = status_counts.get(ComplianceStatus.NON_COMPLIANT, 0)
    not_applicable_count = status_counts.get(ComplianceStatus.NOT_APPLICABLE, 0)
    unmapped_count = status_counts.get(ComplianceStatus.UNMAPPED, 0)
    parameter_queryset = Parameter.objects.all()
    if master and master.parameter_names:
        parameter_queryset = parameter_queryset.filter(name__in=master.parameter_names)
    if readable_ids is not None:
        parameter_queryset = parameter_queryset.filter(pk__in=readable_ids)
    parameter_rows = list(parameter_queryset.annotate(
        state_count=Count("states", filter=Q(states__latest_import__is_current_source=True)),
        non_compliant_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.NON_COMPLIANT)),
        unmapped_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.UNMAPPED)),
    ))
    for parameter in parameter_rows:
        parameter.state_count = states.filter(parameter=parameter).count()

    today = timezone.localdate()
    current_month = today.replace(day=1)
    previous_month = (current_month - timedelta(days=1)).replace(day=1)
    previous_previous_month = (previous_month - timedelta(days=1)).replace(day=1)
    next_month = (current_month.replace(day=28) + timedelta(days=4)).replace(day=1)
    parameter_order = {name.casefold(): index for index, name in enumerate(
        ["Splunk", "CS (Craft)", "RSA", "SNOW", "TGIM", "Logger", "Qualys"]
    )}
    parameter_rows.sort(key=lambda row: (parameter_order.get(row.name.casefold(), 99), row.name.casefold()))

    def month_snapshot(parameter, start, end, label):
        batch = ImportBatch.objects.filter(
            parameter=parameter,
            is_current_source=True,
            snapshot_date__gte=start,
            snapshot_date__lt=end,
            result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
        ).order_by("-snapshot_date", "-imported_at").first()
        if batch is None:
            return {
                "label": label,
                "has_snapshot": False,
                "batch": None,
                "compliant": 0,
                "non_compliant": 0,
                "unmapped": 0,
                "transition_marker_count": 0,
                "compliant_height": 0,
                "non_compliant_height": 0,
                "server_statuses": {},
                "server_names": {},
                "server_ips": {},
                "server_raw_statuses": {},
            }

        records = batch.records.filter(server__isnull=False).exclude(
            issue__contains="Conflicting statuses for this Host Name"
        ).values_list("server_id", "mapped_status", "server__hostname", "raw_status", "ip_address")
        server_statuses = {}
        server_names = {}
        server_ips = {}
        server_raw_statuses = {}
        for server_id, status, hostname, raw_status, record_ip in records:
            if server_id not in server_statuses:
                server_statuses[server_id] = status
                server_names[server_id] = hostname
                server_ips[server_id] = record_ip
                server_raw_statuses[server_id] = raw_status
        return {
            "label": label,
            "has_snapshot": True,
            "batch": batch,
            "compliant": sum(status == ComplianceStatus.COMPLIANT for status in server_statuses.values()),
            "non_compliant": sum(status == ComplianceStatus.NON_COMPLIANT for status in server_statuses.values()),
            "unmapped": sum(status == ComplianceStatus.UNMAPPED for status in server_statuses.values()),
            "transition_marker_count": sum(
                raw_status.strip().casefold() == "compliant to non-compliant"
                for raw_status in server_raw_statuses.values()
            ),
            "compliant_height": 0,
            "non_compliant_height": 0,
            "server_statuses": server_statuses,
            "server_names": server_names,
            "server_ips": server_ips,
            "server_raw_statuses": server_raw_statuses,
        }

    month_windows = [
        (previous_previous_month, previous_month, previous_previous_month.strftime("%b %Y")),
        (previous_month, current_month, previous_month.strftime("%b %Y")),
        (current_month, next_month, current_month.strftime("%b %Y")),
    ]
    parameter_charts = []
    for parameter in parameter_rows:
        periods = [month_snapshot(parameter, start, end, label) for start, end, label in month_windows]
        scale_max = max(
            [period[key] for period in periods for key in ("compliant", "non_compliant")] + [1]
        )
        for period in periods:
            period["compliant_height"] = round(period["compliant"] * 100 / scale_max)
            period["non_compliant_height"] = round(period["non_compliant"] * 100 / scale_max)
        movements = []
        status_labels = dict(ComplianceStatus.choices)
        for previous, current in zip(periods, periods[1:]):
            has_comparison = previous["has_snapshot"] and current["has_snapshot"]
            transition_rows = []
            converted_count = None
            recovered_count = None
            if has_comparison:
                converted_count = sum(
                    previous["server_statuses"].get(server_id) == ComplianceStatus.COMPLIANT
                    and status == ComplianceStatus.NON_COMPLIANT
                    for server_id, status in current["server_statuses"].items()
                )
                recovered_count = sum(
                    previous["server_statuses"].get(server_id) == ComplianceStatus.NON_COMPLIANT
                    and status == ComplianceStatus.COMPLIANT
                    for server_id, status in current["server_statuses"].items()
                )
                for server_id, current_status in current["server_statuses"].items():
                    previous_status = previous["server_statuses"].get(server_id)
                    if (
                        previous_status == ComplianceStatus.COMPLIANT
                        and current_status == ComplianceStatus.NON_COMPLIANT
                    ):
                        transition_rows.append({
                            "hostname": current["server_names"].get(server_id, "Unknown host"),
                            "ip_address": current["server_ips"].get(server_id) or previous["server_ips"].get(server_id, ""),
                            "previous": status_labels.get(previous_status, previous_status),
                            "current": status_labels.get(current_status, current_status),
                        })
                transition_rows.sort(key=lambda row: row["hostname"].casefold())
            movements.append({
                "from_label": previous["label"],
                "to_label": current["label"],
                "has_comparison": has_comparison,
                "converted_count": converted_count,
                "recovered_count": recovered_count,
                "transition_marker_count": current["transition_marker_count"] if current["has_snapshot"] else None,
                "transition_rows": transition_rows,
            })
        parameter_charts.append({
            "id": parameter.pk,
            "name": parameter.name,
            "can_write": has_parameter_permission(request, parameter.pk, "write"),
            "periods": periods,
            "movements": movements,
        })

    monthly_totals = []
    for month_index, (start, end, label) in enumerate(month_windows):
        monthly_totals.append({
            "label": label,
            "compliant": sum(chart["periods"][month_index]["compliant"] for chart in parameter_charts),
            "non_compliant": sum(chart["periods"][month_index]["non_compliant"] for chart in parameter_charts),
            "coverage": sum(chart["periods"][month_index]["has_snapshot"] for chart in parameter_charts),
        })
    monthly_scale = max([
        value for month in monthly_totals
        for value in (month["compliant"], month["non_compliant"])
    ] + [1])
    monthly_comparison = []
    for index, month in enumerate(monthly_totals):
        monthly_comparison.append({
            **month,
            "has_snapshot": month["coverage"] > 0,
            "x": (90, 380, 670)[index],
            "compliant_y": round(184 - month["compliant"] * 145 / monthly_scale),
            "non_compliant_y": round(184 - month["non_compliant"] * 145 / monthly_scale),
        })
    monthly_trend_segments = []
    for previous_month_data, current_month_data in zip(monthly_comparison, monthly_comparison[1:]):
        if not (previous_month_data["has_snapshot"] and current_month_data["has_snapshot"]):
            continue
        for series in ("compliant", "non_compliant"):
            monthly_trend_segments.append({
                "series": series,
                "x1": previous_month_data["x"],
                "y1": previous_month_data[f"{series}_y"],
                "x2": current_month_data["x"],
                "y2": current_month_data[f"{series}_y"],
            })
    monthly_trend_grid = [{
        "value": round(monthly_scale * fraction),
        "y": round(184 - 145 * fraction),
    } for fraction in (1, 0.75, 0.5, 0.25, 0)]
    current_total = compliant_count + non_compliant_count + not_applicable_count + unmapped_count
    applicable_total = compliant_count + non_compliant_count
    compliant_share = round(compliant_count * 100 / applicable_total) if applicable_total else None
    distribution = [
        ("compliant", round(compliant_count * 100 / current_total) if current_total else 0),
        ("non_compliant", round(non_compliant_count * 100 / current_total) if current_total else 0),
        ("not_applicable", round(not_applicable_count * 100 / current_total) if current_total else 0),
        ("unmapped", round(unmapped_count * 100 / current_total) if current_total else 0),
    ]
    status_counts_for_donut = [
        ("#9dd9a9", compliant_count),
        ("#ff9d8f", non_compliant_count),
        ("#c6cfca", not_applicable_count),
        ("#f2cc7e", unmapped_count),
    ]
    donut_stops = []
    if current_total:
        start_percent = 0
        for index, (color, count) in enumerate(status_counts_for_donut):
            end_percent = 100 if index == len(status_counts_for_donut) - 1 else start_percent + count * 100 / current_total
            donut_stops.append(f"{color} {start_percent:.3f}% {end_percent:.3f}%")
            start_percent = end_percent
        status_donut_style = f"conic-gradient({', '.join(donut_stops)})"
    else:
        status_donut_style = "conic-gradient(#dce5de 0% 100%)"

    latest_master_batch = ImportBatch.objects.filter(
        is_current_source=True, is_master_source=True
    ).order_by("-snapshot_date", "-imported_at").first()
    master_server_count = (
        latest_master_batch.records.filter(server__isnull=False)
        .values("server_id").distinct().count()
        if latest_master_batch else 0
    )

    context = {
        "total_servers": master_server_count if master else Server.objects.count(),
        "non_compliant_count": non_compliant_count,
        "compliant_count": compliant_count,
        "not_applicable_count": not_applicable_count,
        "unmapped_count": unmapped_count,
        "recurring_count": states.filter(recurrence_count__gt=0).count(),
        "compliant_share": compliant_share,
        "distribution": distribution,
        "status_donut_style": status_donut_style,
        "parameter_count": len(parameter_rows),
        "parameter_charts": parameter_charts,
        "monthly_comparison": monthly_comparison,
        "monthly_trend_segments": monthly_trend_segments,
        "monthly_trend_grid": monthly_trend_grid,
        "monthly_total_parameters": len(parameter_rows),
        "parameters": parameter_rows,
        "recent_imports": ImportBatch.objects.filter(is_current_source=True, parameter__in=parameter_queryset).select_related("parameter")[:5],
        "recent_events": StatusEvent.objects.filter(
            Q(import_batch__is_current_source=True) | Q(import_batch__isnull=True)
        ).filter(parameter__in=parameter_queryset).select_related("server", "parameter")[:7],
        "master_workbook": master,
    }
    context.update(_dashboard_data_quality(readable_ids, allowed_parameter_ids(request, "write")))
    return render(request, "website/dashboard.html", context)


def new_servers(request):
    master = MasterWorkbook.objects.first()
    details = list(master.new_server_details) if master else []
    if master and master.new_server_count and not details:
        current_batch = master.baseline_batch or ImportBatch.objects.filter(
            is_master_source=True, parameter__name="Splunk"
        ).order_by("-snapshot_date", "-imported_at").first()
        if current_batch:
            previous_batch = ImportBatch.objects.filter(
                is_master_source=True, parameter__name="Splunk"
            ).exclude(pk=current_batch.pk).order_by("-snapshot_date", "-imported_at").first()
            previous_hosts = {
                hostname.strip().casefold()
                for hostname in previous_batch.records.exclude(hostname="").values_list("hostname", flat=True)
            } if previous_batch else set()
            current_details = {}
            for hostname, ip_address in current_batch.records.exclude(hostname="").values_list("hostname", "ip_address"):
                normalized = hostname.strip().casefold()
                if normalized in previous_hosts:
                    continue
                row = current_details.setdefault(normalized, {"hostname": hostname, "ip_addresses": []})
                if ip_address and ip_address not in row["ip_addresses"]:
                    row["ip_addresses"].append(ip_address)
            details = [
                {**row, "ip_address": ", ".join(row["ip_addresses"])}
                for row in sorted(current_details.values(), key=lambda item: item["hostname"].casefold())
            ]

    query = request.GET.get("q", "").strip()
    if query:
        folded_query = query.casefold()
        details = [row for row in details if (
            folded_query in str(row.get("hostname", "")).casefold()
            or folded_query in str(row.get("ip_address", "")).casefold()
        )]
    page, page_size = _paginate(request, details)
    return render(request, "website/new_servers.html", {
        "master_workbook": master,
        "new_server_count": master.new_server_count if master else None,
        "new_servers": page,
        "page_obj": page,
        "page_size": page_size,
        "query": query,
        "total_matches": len(details),
    })


def servers(request):
    search = request.GET.get("q", "").strip()
    state_filter = request.GET.get("status", "")
    environment = request.GET.get("environment", "").strip()
    parameter_filter = request.GET.get("parameter", "")
    master = MasterWorkbook.objects.first()
    parameters = Parameter.objects.all()
    if master and master.parameter_names:
        parameters = parameters.filter(name__in=master.parameter_names)
    readable_ids = allowed_parameter_ids(request, "read")
    if readable_ids is not None:
        parameters = parameters.filter(pk__in=readable_ids)
    selected_parameter = None
    if not parameter_filter and search:
        matching_parameter = parameters.filter(name__iexact=search).first()
        if matching_parameter:
            parameter_filter = str(matching_parameter.pk)
    if parameter_filter.isdigit():
        selected_parameter = parameters.filter(pk=int(parameter_filter)).first()
    current_records = ImportRecord.objects.filter(
        batch__is_current_source=True, server__isnull=False
    )
    if readable_ids is not None:
        current_records = current_records.filter(batch__parameter_id__in=readable_ids)
    current_servers = current_records.values("server_id")
    rows = Server.objects.filter(pk__in=current_servers).annotate(
        non_compliant_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.NON_COMPLIANT, **({"states__parameter_id__in": readable_ids} if readable_ids is not None else {}))),
        unmapped_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.UNMAPPED, **({"states__parameter_id__in": readable_ids} if readable_ids is not None else {}))),
    ).prefetch_related("states__parameter")
    if readable_ids is not None:
        rows = rows.filter(states__latest_import__is_current_source=True, states__parameter_id__in=readable_ids).distinct()
    if search:
        search_filter = Q(hostname__icontains=search) | Q(current_ip__icontains=search)
        if readable_ids is None:
            search_filter |= Q(import_records__batch__is_current_source=True, import_records__ip_address__icontains=search)
            search_filter |= Q(states__latest_import__is_current_source=True, states__parameter__name__icontains=search)
        else:
            search_filter |= Q(import_records__batch__is_current_source=True, import_records__batch__parameter_id__in=readable_ids, import_records__ip_address__icontains=search)
            search_filter |= Q(states__latest_import__is_current_source=True, states__parameter_id__in=readable_ids, states__parameter__name__icontains=search)
        rows = rows.filter(search_filter).distinct()
    if parameter_filter.isdigit():
        rows = rows.filter(
            states__parameter_id=int(parameter_filter),
            states__latest_import__is_current_source=True,
        ).distinct()
    if state_filter in ComplianceStatus.values:
        rows = rows.filter(states__latest_import__is_current_source=True, states__status=state_filter).distinct()
    if environment:
        rows = rows.filter(environment__iexact=environment)
    page, page_size = _paginate(request, rows)
    environments = Server.objects.filter(pk__in=current_servers).exclude(environment="").values_list("environment", flat=True).distinct().order_by("environment")
    if selected_parameter:
        selected_states = ComplianceState.objects.filter(
            server_id__in=[server.pk for server in page.object_list],
            parameter=selected_parameter,
            latest_import__is_current_source=True,
        ).select_related("parameter")
        states_by_server = {state.server_id: state for state in selected_states}
        for server in page.object_list:
            server.parameter_state = states_by_server.get(server.pk)
    return render(request, "website/servers.html", {
        "page_obj": page,
        "page_size": page_size,
        "search": search,
        "status_filter": state_filter,
        "environment_filter": environment,
        "parameter_filter": parameter_filter,
        "selected_parameter": selected_parameter,
        "parameters": parameters,
        "environments": environments,
        "status_choices": ComplianceStatus.choices,
    })


def non_compliant(request):
    search = request.GET.get("q", "").strip()
    parameter_id = request.GET.get("parameter", "")
    rows = ComplianceState.objects.filter(
        status=ComplianceStatus.NON_COMPLIANT,
        latest_import__is_current_source=True,
    ).select_related("server", "parameter", "latest_import")
    readable_ids = allowed_parameter_ids(request, "read")
    if readable_ids is not None:
        rows = rows.filter(parameter_id__in=readable_ids)
    if search:
        search_filter = (
            Q(server__hostname__icontains=search)
            | Q(server__current_ip__icontains=search)
        )
        if readable_ids is None:
            search_filter |= Q(server__import_records__ip_address__icontains=search)
        else:
            search_filter |= Q(server__import_records__batch__parameter_id__in=readable_ids, server__import_records__batch__is_current_source=True, server__import_records__ip_address__icontains=search)
        rows = rows.filter(search_filter).distinct()
    if parameter_id.isdigit():
        rows = rows.filter(parameter_id=int(parameter_id))
    page, page_size = _paginate(request, rows)
    return render(request, "website/non_compliant.html", {
        "page_obj": page,
        "page_size": page_size,
        "search": search,
        "parameter_filter": parameter_id,
        "parameters": Parameter.objects.filter(
            name__in=(MasterWorkbook.objects.first().parameter_names if MasterWorkbook.objects.first() else Parameter.objects.values_list("name", flat=True))
        ).filter(pk__in=readable_ids) if readable_ids is not None else Parameter.objects.filter(
            name__in=(MasterWorkbook.objects.first().parameter_names if MasterWorkbook.objects.first() else Parameter.objects.values_list("name", flat=True))
        ),
    })


def server_detail(request, server_id):
    readable_ids = allowed_parameter_ids(request, "read")
    readable_records = ImportRecord.objects.filter(batch__is_current_source=True)
    if readable_ids is not None:
        readable_records = readable_records.filter(batch__parameter_id__in=readable_ids)
    server = get_object_or_404(
        Server.objects.filter(pk__in=readable_records.values("server_id")).distinct(), pk=server_id
    )
    change_form = ManualStateChangeForm(request.POST or None)
    master = MasterWorkbook.objects.first()
    writable_ids = allowed_parameter_ids(request, "write")
    if master and master.parameter_names:
        change_form.fields["parameter"].queryset = Parameter.objects.filter(name__in=master.parameter_names)
    master_parameter_ids = ImportBatch.objects.filter(
        is_current_source=True, is_master_source=True
    ).values_list("parameter_id", flat=True).distinct()
    has_master_write_access = request.workspace_is_owner or any(
        has_parameter_permission(request, parameter_id, "write") for parameter_id in master_parameter_ids
    )
    change_form.fields["parameter"].queryset = change_form.fields["parameter"].queryset.exclude(pk__in=master_parameter_ids)
    if writable_ids is not None:
        change_form.fields["parameter"].queryset = change_form.fields["parameter"].queryset.filter(pk__in=writable_ids)
    if request.method == "POST" and change_form.is_valid():
        parameter = change_form.cleaned_data["parameter"]
        new_status = change_form.cleaned_data["status"]
        reason = change_form.cleaned_data["reason"].strip()
        event_date = timezone.localdate()
        state = ComplianceState.objects.filter(server=server, parameter=parameter).filter(
            Q(latest_import__is_current_source=True) | Q(latest_import__isnull=True)
        ).first()
        previous_status = state.status if state else ""
        if state and state.status == new_status:
            AuditLog.objects.create(
                action="manual_status_no_change",
                detail=f"{parameter.name}: {state.get_status_display()}. Note: {reason}",
                server=server,
            )
            messages.info(request, "The current status already matches. Your note was added to the audit log.")
        else:
            recurrent = bool(
                state
                and previous_status == ComplianceStatus.COMPLIANT
                and new_status == ComplianceStatus.NON_COMPLIANT
                and StatusEvent.objects.filter(
                    server=server,
                    parameter=parameter,
                    new_status=ComplianceStatus.NON_COMPLIANT,
                ).filter(Q(import_batch__is_current_source=True) | Q(import_batch__isnull=True)).exists()
            )
            if state is None:
                state = ComplianceState(
                    server=server,
                    parameter=parameter,
                    observed_on=event_date,
                )
                kind = StatusEvent.Kind.MANUAL
            else:
                kind = StatusEvent.Kind.RECURRENCE if recurrent else StatusEvent.Kind.MANUAL
            state.status = new_status
            state.raw_status = "Manual update"
            state.observed_on = event_date
            if recurrent:
                state.recurrence_count += 1
            state.save()
            StatusEvent.objects.create(
                server=server,
                parameter=parameter,
                previous_status=previous_status,
                new_status=new_status,
                kind=kind,
                event_date=event_date,
                reason=reason,
            )
            AuditLog.objects.create(
                action="manual_status_update",
                detail=f"{parameter.name}: {previous_status or 'No prior status'} → {new_status}. {reason}",
                server=server,
            )
            messages.success(request, "The current status was updated and a history event was recorded.")
        return redirect("server_detail", server_id=server.id)

    states = server.states.filter(
        Q(latest_import__is_current_source=True) | Q(latest_import__isnull=True)
    ).select_related("parameter", "latest_import")
    events = server.status_events.filter(
        Q(import_batch__is_current_source=True) | Q(import_batch__isnull=True)
    ).select_related("parameter", "import_batch")
    source_rows = server.import_records.filter(batch__is_current_source=True)
    if readable_ids is not None:
        source_rows = source_rows.filter(batch__parameter_id__in=readable_ids)
    import_records = list(source_rows.select_related("batch", "batch__parameter").order_by(
        "-batch__snapshot_date", "-batch__imported_at", "source_row"
    )[:20])
    if readable_ids is not None:
        states = states.filter(parameter_id__in=readable_ids)
        events = events.filter(parameter_id__in=readable_ids)
    events = events[:30]
    latest_batch_ids = {}
    for parameter_id in {record.batch.parameter_id for record in import_records}:
        latest = ImportBatch.objects.filter(
            parameter_id=parameter_id,
            is_current_source=True,
            result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
        ).order_by("-snapshot_date", "-imported_at").first()
        latest_batch_ids[parameter_id] = latest.pk if latest else None
    for record in import_records:
        record.can_edit = (
            latest_batch_ids.get(record.batch.parameter_id) == record.batch_id
            and has_parameter_permission(request, record.batch.parameter_id, "write")
        )
    editable_records = [record for record in import_records if record.can_edit]
    return render(request, "website/server_detail.html", {
        "server": server,
        "states": states,
        "events": events,
        "import_records": import_records,
        "editable_records": editable_records,
        "change_form": change_form,
        "master_workbook_active": bool(master and master.path),
        "has_master_write_access": has_master_write_access,
    })


def parameters(request):
    master = MasterWorkbook.objects.first()
    parameter_rows = Parameter.objects.all()
    if master and master.parameter_names:
        parameter_rows = parameter_rows.filter(name__in=master.parameter_names)
    readable_ids = allowed_parameter_ids(request, "read")
    if readable_ids is not None:
        parameter_rows = parameter_rows.filter(pk__in=readable_ids)
    rows = parameter_rows.annotate(
        server_count=Count("states", filter=Q(states__latest_import__is_current_source=True)),
        non_compliant_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.NON_COMPLIANT)),
        unmapped_count=Count("states", filter=Q(states__latest_import__is_current_source=True, states__status=ComplianceStatus.UNMAPPED)),
        latest_snapshot_date=Max("imports__snapshot_date", filter=Q(imports__is_current_source=True)),
    )
    for parameter in rows:
        parameter.can_write = has_parameter_permission(request, parameter.pk, "write")
        if master:
            parameter.server_count = master.roster_count
    return render(request, "website/parameters.html", {"parameters": rows})


def parameter_data(request, parameter_id):
    parameter = get_object_or_404(Parameter, pk=parameter_id)
    if not has_parameter_permission(request, parameter.pk, "read"):
        return HttpResponseForbidden("You do not have Read access to this parameter.")
    try:
        dataset = read_master_parameter_sheet(parameter.name)
    except ImportValidationError as exc:
        messages.error(request, str(exc))
        return redirect("parameters")
    latest_batch = ImportBatch.objects.filter(
        parameter=parameter, is_current_source=True, is_master_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    record_ids = {}
    if latest_batch:
        record_ids = dict(latest_batch.records.values_list("source_row", "id"))
    query = request.GET.get("q", "").strip().casefold()
    data_rows = []
    can_write = has_parameter_permission(request, parameter.pk, "write")
    for raw_row in dataset["rows"]:
        row = {
            **raw_row,
            "edit_url": "",
            "cells": [
                {
                    "heading": heading,
                    "value": raw_row["values"].get(heading),
                    "has_value": raw_row["values"].get(heading) not in (None, ""),
                }
                for heading in dataset["headers"]
            ],
        }
        if can_write:
            record_id = record_ids.get(raw_row["source_row"])
            row["edit_url"] = reverse("source_record_edit", args=[record_id]) if record_id else reverse(
                "master_sheet_row_edit", args=[parameter.pk, raw_row["source_row"]]
            )
        haystack = " ".join("" if value is None else str(value) for value in raw_row["values"].values()).casefold()
        if not query or query in haystack:
            data_rows.append(row)
    page, page_size = _paginate(request, data_rows)
    return render(request, "website/parameter_data.html", {
        "parameter": parameter,
        "sheet_name": dataset["sheet_name"],
        "headers": dataset["headers"],
        "page_obj": page,
        "page_size": page_size,
        "search": request.GET.get("q", "").strip(),
        "can_write": can_write,
        "snapshot": latest_batch,
        "table_colspan": len(dataset["headers"]) + 1 + (1 if can_write else 0),
    })


def master_sheet_row_edit(request, parameter_id, source_row):
    parameter = get_object_or_404(Parameter, pk=parameter_id)
    if not has_parameter_permission(request, parameter.pk, "write"):
        return HttpResponseForbidden("You do not have Write access to this parameter.")
    try:
        dataset = read_master_parameter_sheet(parameter.name)
    except ImportValidationError as exc:
        messages.error(request, str(exc))
        return redirect("parameter_data", parameter_id=parameter.pk)
    row = next((item for item in dataset["rows"] if item["source_row"] == source_row), None)
    if row is None:
        raise Http404("That worksheet row was not found.")
    latest_batch = ImportBatch.objects.filter(
        parameter=parameter, is_current_source=True, is_master_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    if latest_batch:
        record_id = latest_batch.records.filter(source_row=source_row).values_list("id", flat=True).first()
        if record_id:
            return redirect("source_record_edit", record_id=record_id)
    form = MasterSheetRowEditForm(
        row, dataset["headers"], row.get("formula_columns", ()),
        request.POST or None, actor=request.workspace_user_id,
    )
    if request.method == "POST" and form.is_valid():
        updates = form.changed_columns()
        try:
            edit_master_sheet_row(parameter.name, source_row, updates, request.workspace_user_id)
        except ImportValidationError as exc:
            form.add_error(None, str(exc))
        except OSError:
            form.add_error(None, "The workbook could not be saved. Close it in Excel and try again.")
        else:
            messages.success(request, "The row was saved to the master workbook and the dashboard snapshots were refreshed.")
            return redirect("parameter_data", parameter_id=parameter.pk)
    return render(request, "website/master_sheet_row_edit.html", {
        "parameter": parameter,
        "sheet_name": dataset["sheet_name"],
        "row": row,
        "form": form,
    })


def compliance_summary(request):
    master = MasterWorkbook.objects.first()
    readable_ids = allowed_parameter_ids(request, "read")
    parameters = Parameter.objects.all()
    if master and master.parameter_names:
        parameters = parameters.filter(name__in=master.parameter_names)
    if readable_ids is not None:
        parameters = parameters.filter(pk__in=readable_ids)
    parameters = list(parameters)
    parameter_positions = {name: index for index, name in enumerate(master.parameter_names)} if master else {}
    parameters.sort(key=lambda item: (parameter_positions.get(item.name, 10_000), item.name.casefold()))
    for parameter in parameters:
        parameter.display_name = parameter_sheet_name(parameter.name)
    roster_batch = ImportBatch.objects.filter(
        is_master_source=True, is_current_source=True, parameter__name="Splunk"
    ).order_by("-snapshot_date", "-imported_at").first()
    if roster_batch:
        server_ids = list(roster_batch.records.exclude(server_id=None).values_list("server_id", flat=True).distinct())
        servers_list = list(Server.objects.filter(pk__in=server_ids).order_by("hostname"))
    else:
        state_server_ids = ComplianceState.objects.filter(latest_import__is_current_source=True)
        if readable_ids is not None:
            state_server_ids = state_server_ids.filter(parameter_id__in=readable_ids)
        servers_list = list(Server.objects.filter(states__in=state_server_ids).distinct().order_by("hostname"))
    state_query = ComplianceState.objects.filter(latest_import__is_current_source=True, parameter__in=parameters).select_related("parameter")
    states_by_pair = {(state.server_id, state.parameter_id): state for state in state_query}
    query = request.GET.get("q", "").strip().casefold()
    status_filter = request.GET.get("status", "")
    summary_rows = []
    compliant_total = non_compliant_total = pending_total = 0
    for server in servers_list:
        param_cells = []
        mapped_compliant = mapped_non_compliant = 0
        updates = []
        for parameter in parameters:
            state = states_by_pair.get((server.pk, parameter.pk))
            if state:
                label = state.get_status_display() if state.status != ComplianceStatus.UNMAPPED else (state.raw_status or "Unmapped")
                status = state.status
                updates.append(state.updated_at)
            else:
                label, status = "Not Started", "pending"
            if status in {ComplianceStatus.COMPLIANT, ComplianceStatus.NOT_APPLICABLE}:
                mapped_compliant += 1
            elif status == ComplianceStatus.NON_COMPLIANT:
                mapped_non_compliant += 1
            param_cells.append({"name": parameter_sheet_name(parameter.name), "label": label, "status": status})
        pending = max(0, len(parameters) - mapped_compliant - mapped_non_compliant)
        if mapped_non_compliant:
            overall = "Non-Compliant"
            overall_key = "non_compliant"
        elif parameters and mapped_compliant == len(parameters):
            overall = "Compliant"
            overall_key = "compliant"
        else:
            overall = "Pending"
            overall_key = "pending"
        if query and query not in server.hostname.casefold() and query not in server.current_ip.casefold():
            continue
        if status_filter and status_filter != overall_key:
            continue
        row = {
            "server": server,
            "cells": param_cells,
            "compliant_count": mapped_compliant,
            "non_compliant_count": mapped_non_compliant,
            "pending_count": pending,
            "overall": overall,
            "overall_key": overall_key,
            "last_updated": max(updates) if updates else None,
            "percent": round(mapped_compliant * 100 / len(parameters)) if parameters else 0,
        }
        summary_rows.append(row)
    compliant_total = sum(row["compliant_count"] for row in summary_rows)
    non_compliant_total = sum(row["non_compliant_count"] for row in summary_rows)
    pending_total = sum(row["pending_count"] for row in summary_rows)
    page, page_size = _paginate(request, summary_rows)
    return render(request, "website/compliance_summary.html", {
        "parameters": parameters,
        "page_obj": page,
        "page_size": page_size,
        "search": request.GET.get("q", "").strip(),
        "status_filter": status_filter,
        "compliant_total": compliant_total,
        "non_compliant_total": non_compliant_total,
        "pending_total": pending_total,
        "summary_count": len(summary_rows),
    })


def imports(request):
    search = request.GET.get("q", "").strip()
    snapshot_filter = request.GET.get("snapshot_date", "").strip()
    uploaded_filter = request.GET.get("uploaded_on", "").strip()
    parameter_filter = request.GET.get("parameter", "")
    batches = ImportBatch.objects.select_related("parameter")
    readable_ids = allowed_parameter_ids(request, "read")
    if readable_ids is not None:
        batches = batches.filter(parameter_id__in=readable_ids)
    if search:
        batches = batches.filter(
            Q(parameter__name__icontains=search) | Q(original_filename__icontains=search)
        )
    parsed_snapshot = parse_date(snapshot_filter) if snapshot_filter else None
    parsed_uploaded = parse_date(uploaded_filter) if uploaded_filter else None
    if parsed_snapshot:
        batches = batches.filter(snapshot_date=parsed_snapshot)
    if parsed_uploaded:
        batches = batches.filter(imported_at__date=parsed_uploaded)
    if parameter_filter.isdigit():
        batches = batches.filter(parameter_id=int(parameter_filter))
    page, page_size = _paginate(request, batches)
    return render(request, "website/imports.html", {
        "page_obj": page,
        "page_size": page_size,
        "search": search,
        "snapshot_filter": snapshot_filter,
        "uploaded_filter": uploaded_filter,
        "parameter_filter": parameter_filter,
        "parameters": Parameter.objects.filter(pk__in=readable_ids) if readable_ids is not None else Parameter.objects.all(),
    })


def delete_import(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        result = delete_import_batch(batch_id, actor_id=request.workspace_user_id)
    except ImportBatch.DoesNotExist as exc:
        raise Http404("That import was already deleted.") from exc
    messages.success(
        request,
        f"Deleted {result['parameter']} snapshot dated {result['snapshot_date']:%d %b %Y}. "
        "The original workbook was not changed.",
    )
    return redirect("imports")


def delete_history_event(request, event_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        delete_manual_status_event(event_id, actor_id=request.workspace_user_id)
    except StatusEvent.DoesNotExist as exc:
        raise Http404("That history event was already deleted.") from exc
    except ValueError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "The manual history event was deleted.")
    return redirect("history")


def clear_all_data(request):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    if request.POST.get("confirmation", "").strip() != "DELETE ALL DATA":
        messages.error(request, 'Type "DELETE ALL DATA" exactly to confirm the clear operation.')
        return redirect("history")
    result = clear_compliance_workspace_data(actor_id=request.workspace_user_id)
    messages.success(
        request,
        f"Cleared {result['imports']} import(s), {result['servers']} server(s), and "
        f"{result['events']} status event(s). User access accounts and original Excel files were kept.",
    )
    return redirect("history")


def imports_new(request):
    form = ExcelImportForm(request.POST or None, request.FILES or None)
    if not request.workspace_is_owner:
        form.fields.pop("sync_target_path", None)
    if request.method == "POST" and form.is_valid():
        uploaded = form.cleaned_data["excel_file"]
        content = uploaded.read()
        try:
            if not request.workspace_is_owner:
                requested_parameter = Parameter.objects.filter(name__iexact=form.cleaned_data["parameter_name"].strip()).first()
                if requested_parameter is None or not has_parameter_permission(request, requested_parameter.pk, "execute"):
                    raise ImportValidationError("You do not have Execute access for this parameter. Ask the owner for a grant.")
            preview = preview_workbook(content, form.cleaned_data["status_column"])
            snapshot_date = form.cleaned_data["snapshot_date"] or preview["suggested_date"]
            if snapshot_date is None:
                raise ImportValidationError(
                    "No date was found in the status heading. Enter the snapshot date and import again."
                )
            if not request.workspace_is_owner:
                parameter = Parameter.objects.filter(name=form.cleaned_data["parameter_name"]).first()
                latest_schema = None
                if parameter:
                    latest_schema = ImportBatch.objects.filter(
                        parameter=parameter,
                        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
                    ).order_by("-snapshot_date", "-imported_at").first()
                if latest_schema is None:
                    raise ImportValidationError("Only the owner can establish a parameter's first workbook schema.")
                expected_columns = [
                    heading for heading in latest_schema.source_columns
                    if heading != latest_schema.status_column
                ]
                incoming_columns = [
                    heading for heading in preview["headers"]
                    if heading != preview["status_column"]
                ]
                if incoming_columns != expected_columns:
                    raise ImportValidationError(
                        "This workbook changes the existing column names, number, or order. Ask the owner to approve the new schema first."
                    )
            sync_target_path = validate_sync_target_path(
                form.cleaned_data.get("sync_target_path", "") if request.workspace_is_owner else "",
                expected_hash=sha256(content).hexdigest(),
                source_sheet=preview["source_sheet"],
                source_columns=preview["headers"],
            )
            batch = import_snapshot(
                content,
                original_filename=uploaded.name,
                parameter_name=form.cleaned_data["parameter_name"],
                snapshot_date=snapshot_date,
                requested_status_column=form.cleaned_data["status_column"],
                preview=preview,
                sync_target_path=sync_target_path,
            )
        except ImportValidationError as exc:
            AuditLog.objects.create(action="excel_import_failed", detail=str(exc)[:2000])
            form.add_error(None, str(exc))
        except Exception:
            AuditLog.objects.create(
                action="excel_import_failed",
                detail="The upload could not be imported due to an unexpected processing error.",
            )
            form.add_error(None, "The workbook could not be imported. Check its columns and try again.")
        else:
            messages.success(request, f"Imported {batch.total_rows:,} rows for {batch.parameter.name}.")
            return redirect("import_detail", batch_id=batch.id)
    return render(request, "website/import_new.html", {"form": form})


def imports_master(request):
    config = MasterWorkbook.objects.first()
    initial = {"target_path": config.path} if config and config.path else None
    form = MasterWorkbookImportForm(request.POST or None, request.FILES or None, initial=initial)
    if not request.workspace_is_owner:
        form.fields.pop("target_path", None)
    if request.method == "POST" and form.is_valid():
        uploaded = form.cleaned_data["master_file"]
        content = uploaded.read()
        if request.workspace_is_owner:
            target_path = form.cleaned_data.get("target_path", "")
        else:
            target_path = config.path if config else ""
        if not request.workspace_is_owner and not target_path:
            form.add_error(None, "Ask the owner to configure the master workbook write-back path first.")
        else:
            try:
                result = import_master_workbook(
                    content,
                    original_filename=uploaded.name,
                    target_path=target_path,
                )
            except ImportValidationError as exc:
                AuditLog.objects.create(action="master_workbook_import_failed", detail=str(exc)[:2000])
                form.add_error(None, str(exc))
            except Exception:
                AuditLog.objects.create(
                    action="master_workbook_import_failed",
                    detail="The master workbook could not be imported due to an unexpected processing error.",
                )
                form.add_error(None, "The master workbook could not be imported. Check its sheets and headings, then try again.")
            else:
                imported_count = len(result["imported"])
                if result["new_server_count"] is None:
                    detail = (
                        f"Set the baseline from {result['roster_count']:,} server hostnames. "
                        "New-server counts will be available with the next master workbook upload."
                    )
                else:
                    detail = f"{result['new_server_count']:,} new server(s) found against the previous master roster."
                messages.success(request, f"Master workbook saved. Imported {imported_count} new or changed monthly column(s). {detail}")
                return redirect("dashboard")
    return render(request, "website/master_import.html", {"form": form, "master_workbook": config})


def import_detail(request, batch_id):
    batch = get_object_or_404(ImportBatch.objects.select_related("parameter"), pk=batch_id)
    if not has_parameter_permission(request, batch.parameter_id, "read"):
        return HttpResponseForbidden("You do not have Read access to this parameter.")
    counts = list(
        batch.records.values("raw_status", "mapped_status")
        .annotate(row_count=Count("id"))
        .order_by("raw_status")
    )
    issues = batch.records.filter(Q(issue__gt="") | Q(duplicate_row=True)).select_related("server")
    issue_page, page_size = _paginate(request, issues)
    latest_batch = ImportBatch.objects.filter(
        parameter=batch.parameter,
        is_current_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    return render(request, "website/import_detail.html", {
        "batch": batch,
        "status_counts": counts,
        "issue_page": issue_page,
        "page_size": page_size,
        "can_edit_rows": bool(latest_batch and latest_batch.pk == batch.pk and has_parameter_permission(request, batch.parameter_id, "write")),
        "can_download_workbook": request.workspace_is_owner or not batch.is_master_source,
        "is_master_workbook": batch.is_master_source,
        "sync_target_form": SyncTargetForm(initial={"sync_target_path": batch.sync_target_path}),
    })


def download_workbook(request, batch_id):
    batch = get_object_or_404(ImportBatch, pk=batch_id)
    if not has_parameter_permission(request, batch.parameter_id, "read"):
        return HttpResponseForbidden("You do not have Read access to this parameter.")
    if batch.is_master_source and not request.workspace_is_owner:
        return HttpResponseForbidden("Only the owner can download the full master workbook. Open your permitted worksheet in the app instead.")
    workbook_file = batch.working_file or batch.source_file
    if not workbook_file:
        raise Http404("No workbook is available for this import.")
    try:
        file_handle = workbook_file.open("rb")
    except (FileNotFoundError, OSError) as exc:
        raise Http404("The saved workbook is no longer available.") from exc
    filename = f"{slugify(Path(batch.original_filename).stem) or 'workbook'}_current.xlsx"
    response = FileResponse(
        file_handle,
        as_attachment=True,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Cache-Control"] = "private, no-store"
    return response


def set_import_sync_target(request, batch_id):
    batch = get_object_or_404(ImportBatch.objects.select_related("parameter"), pk=batch_id)
    if batch.is_master_source:
        messages.error(request, "Use the master workbook import page to configure the shared write-back path.")
        return redirect("import_detail", batch_id=batch.id)
    if request.method != "POST":
        return redirect("import_detail", batch_id=batch.id)
    form = SyncTargetForm(request.POST)
    if form.is_valid():
        try:
            target_path = validate_sync_target_path(
                form.cleaned_data["sync_target_path"],
                expected_hash=batch.working_sha256 or batch.file_sha256,
                source_sheet=batch.source_sheet,
                source_columns=batch.source_columns,
            )
        except ImportValidationError as exc:
            messages.error(request, str(exc))
        else:
            batch.sync_target_path = target_path
            batch.save(update_fields=["sync_target_path"])
            AuditLog.objects.create(
                action="workbook_sync_target_set" if target_path else "workbook_sync_target_cleared",
                detail=("Original workbook write-back enabled." if target_path else "Original workbook write-back disabled."),
                import_batch=batch,
            )
            messages.success(
                request,
                "Web edits will update the matching original workbook." if target_path
                else "Web edits will update the workbook copy stored by this app.",
            )
    else:
        messages.error(request, "Enter a valid full path to the matching workbook.")
    return redirect("import_detail", batch_id=batch.id)


def source_record_edit(request, record_id):
    record = get_object_or_404(
        ImportRecord.objects.select_related("batch", "batch__parameter", "server"),
        pk=record_id,
    )
    if not has_parameter_permission(request, record.batch.parameter_id, "write"):
        return HttpResponseForbidden("You do not have Write access to this parameter.")
    latest_batch = ImportBatch.objects.filter(
        parameter=record.batch.parameter,
        is_current_source=True,
        result__in=[ImportBatch.Result.COMPLETED, ImportBatch.Result.PARTIAL],
    ).order_by("-snapshot_date", "-imported_at").first()
    if latest_batch is None or latest_batch.pk != record.batch_id:
        messages.error(request, "Only rows in the latest snapshot can be edited. Re-import a newer workbook first.")
        return redirect("import_detail", batch_id=record.batch_id)
    form = SourceRecordEditForm(record, request.POST or None, actor=request.workspace_user_id)
    if request.method == "POST" and form.is_valid():
        actor = request.workspace_user_id
        updates = form.changed_columns()
        try:
            updated_record = edit_source_record(record.id, updates, actor)
        except ImportValidationError as exc:
            form.add_error(None, str(exc))
        except OSError:
            form.add_error(None, "The workbook could not be saved. Close it in Excel and try again.")
        else:
            messages.success(request, "The row and the master workbook were updated." if record.batch.is_master_source else ("The row and its workbook copy were updated." if not record.batch.sync_target_path else "The row and original workbook were updated."))
            if updated_record.server_id:
                return redirect("server_detail", server_id=updated_record.server_id)
            return redirect("import_detail", batch_id=updated_record.batch_id)
    return render(request, "website/source_record_edit.html", {
        "record": record,
        "batch": record.batch,
        "form": form,
    })


def history(request):
    search = request.GET.get("q", "").strip()
    parameter_id = request.GET.get("parameter", "")
    recurring_only = request.GET.get("recurring", "") == "1"
    rows = StatusEvent.objects.filter(
        Q(import_batch__is_current_source=True) | Q(import_batch__isnull=True)
    ).select_related("server", "parameter", "import_batch")
    readable_ids = allowed_parameter_ids(request, "read")
    if readable_ids is not None:
        rows = rows.filter(parameter_id__in=readable_ids)
    if search:
        search_filter = (
            Q(server__hostname__icontains=search)
            | Q(server__current_ip__icontains=search)
        )
        if readable_ids is None:
            search_filter |= Q(server__import_records__ip_address__icontains=search)
        else:
            search_filter |= Q(server__import_records__batch__parameter_id__in=readable_ids, server__import_records__ip_address__icontains=search)
        rows = rows.filter(search_filter).distinct()
    if parameter_id.isdigit():
        rows = rows.filter(parameter_id=int(parameter_id))
    if recurring_only:
        rows = rows.filter(kind=StatusEvent.Kind.RECURRENCE)
    page, page_size = _paginate(request, rows)
    return render(request, "website/history.html", {
        "page_obj": page,
        "page_size": page_size,
        "search": search,
        "parameter_filter": parameter_id,
        "recurring_only": recurring_only,
        "parameters": Parameter.objects.filter(
            name__in=(MasterWorkbook.objects.first().parameter_names if MasterWorkbook.objects.first() else Parameter.objects.values_list("name", flat=True))
        ).filter(pk__in=readable_ids) if readable_ids is not None else Parameter.objects.filter(
            name__in=(MasterWorkbook.objects.first().parameter_names if MasterWorkbook.objects.first() else Parameter.objects.values_list("name", flat=True))
        ),
    })


def export_non_compliant(request):
    search = request.GET.get("q", "").strip()
    parameter_id = request.GET.get("parameter", "")
    rows = ComplianceState.objects.filter(
        status=ComplianceStatus.NON_COMPLIANT,
        latest_import__is_current_source=True,
    ).select_related("server", "parameter", "latest_import")
    executable_ids = allowed_parameter_ids(request, "execute")
    if executable_ids is not None:
        rows = rows.filter(parameter_id__in=executable_ids)
    if search:
        search_filter = (
            Q(server__hostname__icontains=search)
            | Q(server__current_ip__icontains=search)
        )
        if executable_ids is None:
            search_filter |= Q(server__import_records__ip_address__icontains=search)
        else:
            search_filter |= Q(server__import_records__batch__parameter_id__in=executable_ids, server__import_records__ip_address__icontains=search)
        rows = rows.filter(search_filter).distinct()
    if parameter_id.isdigit():
        rows = rows.filter(parameter_id=int(parameter_id))

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Current non-compliant")
    sheet.append([
        "Host Name", "IP Address", "OPERATING SYSTEM", "Environment", "Parameter",
        "Current status", "Source status label", "First observed/current date", "Recurrence count",
    ])
    for state in rows.iterator(chunk_size=500):
        values = [
            state.server.hostname,
            state.server.current_ip,
            state.server.operating_system,
            state.server.environment,
            state.parameter.name,
            state.get_status_display(),
            state.raw_status,
            state.observed_on.isoformat(),
            state.recurrence_count,
        ]
        safe_cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            safe_cells.append(cell)
        sheet.append(safe_cells)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="current_non_compliant_servers.xlsx"'
    return response
